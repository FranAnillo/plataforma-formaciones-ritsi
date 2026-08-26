from __future__ import annotations

import io
import asyncio
import base64
import hashlib
import html as html_lib
import hmac
import logging
import os
import re
import secrets
import unicodedata
import uuid
from contextlib import asynccontextmanager
from datetime import date, datetime, timedelta, timezone
from enum import Enum
from pathlib import Path
from typing import Any, Dict, List, Optional
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError

import bcrypt
import requests
from cryptography.fernet import Fernet, InvalidToken
from dotenv import load_dotenv
from fastapi import APIRouter, Cookie, Depends, FastAPI, HTTPException, Request
from fastapi.responses import JSONResponse, PlainTextResponse, RedirectResponse
from motor.motor_asyncio import AsyncIOMotorClient
from openpyxl import load_workbook
from pymongo import ReturnDocument
from pydantic import BaseModel, ConfigDict, EmailStr, Field, field_validator, model_validator
from starlette.middleware.cors import CORSMiddleware

try:
    from .integrations import render_ics, telegram_link
except ImportError:
    from integrations import render_ics, telegram_link

ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / ".env")
load_dotenv(ROOT_DIR.parent / ".env")
MONGO_URL = os.getenv("MONGO_URL", "mongodb://localhost:27017")
DB_NAME = os.getenv("DB_NAME", "plataforma_formativa_ritsi")
SHEET_ID = os.getenv("FORMATIONS_SHEET_ID", "1JSRrepNNdQDl6zeroZPDKUJztLhSC3j6r7y2h_ND09c")
SESSION_DAYS = int(os.getenv("SESSION_DAYS", "7"))
COOKIE_SECURE = os.getenv("COOKIE_SECURE", "false").lower() == "true"
MEETING_URL_ENCRYPTION_KEY = os.getenv("MEETING_URL_ENCRYPTION_KEY")
DEFAULT_TIMEZONE = os.getenv("DEFAULT_TIMEZONE", "Europe/Madrid")
DEFAULT_JOIN_WINDOW_MINUTES = int(os.getenv("DEFAULT_JOIN_WINDOW_MINUTES", "15"))
PUBLIC_APP_URL = os.getenv("PUBLIC_APP_URL", "http://localhost:3000")
PUBLIC_API_URL = os.getenv("PUBLIC_API_URL", "http://localhost:8000")
ALLOWED_ORIGINS = {
    value.strip()
    for value in os.getenv("CORS_ORIGINS", "http://localhost:3000").split(",")
    if value.strip()
}
UNSAFE_HTTP_METHODS = {"POST", "PUT", "PATCH", "DELETE"}

client = AsyncIOMotorClient(MONGO_URL)
db = client[DB_NAME]
api_router = APIRouter(prefix="/api")
logger = logging.getLogger(__name__)


class UserType(str, Enum):
    REPRESENTANTE = "representante"
    UNIVERSIDAD = "universidad"
    JUNTA_DIRECTIVA = "junta_directiva"
    FORMADOR = "formador"
    COLABORACION_EXTERNA = "colaboracion_externa"
    ADMIN = "admin"


class BoardPosition(str, Enum):
    PRESIDENCIA = "presidencia"
    VICEPRESIDENCIA_POLITICA_UNIVERSITARIA = "vicepresidencia_politica_universitaria"
    TESORERIA = "tesoreria"
    SECRETARIA = "secretaria"
    VICEPRESIDENCIA_COMUNICACION = "vicepresidencia_comunicacion"
    MIEMBRO_ADICIONAL_1 = "miembro_adicional_1"
    MIEMBRO_ADICIONAL_2 = "miembro_adicional_2"


REQUIRED_BOARD_POSITIONS = {
    BoardPosition.PRESIDENCIA,
    BoardPosition.VICEPRESIDENCIA_POLITICA_UNIVERSITARIA,
    BoardPosition.TESORERIA,
    BoardPosition.SECRETARIA,
    BoardPosition.VICEPRESIDENCIA_COMUNICACION,
}


class ContentStatus(str, Enum):
    DRAFT = "draft"
    PENDING = "pending"
    CHANGES_REQUESTED = "changes_requested"
    PUBLISHED = "published"
    ARCHIVED = "archived"


class EnrollmentStatus(str, Enum):
    ASSIGNED = "assigned"
    IN_PROGRESS = "in_progress"
    COMPLETED = "completed"
    CANCELLED = "cancelled"


class EnrollmentSource(str, Enum):
    ASSIGNMENT = "assignment"
    GLOBAL_ASSIGNMENT = "global_assignment"
    SELF_ENROLLED = "self_enrolled"


class SynchronousSessionStatus(str, Enum):
    DRAFT = "draft"
    SCHEDULED = "scheduled"
    LIVE = "live"
    ENDED = "ended"
    CANCELLED = "cancelled"


class FileType(str, Enum):
    VIDEO = "video"
    PDF = "pdf"
    IMAGE = "image"
    PRESENTATION = "presentation"
    DOCUMENT = "document"
    LINK = "link"


class QuestionType(str, Enum):
    TRUE_FALSE = "true_false"
    MULTIPLE_CHOICE = "multiple_choice"
    MULTIPLE_RESPONSE = "multiple_response"


class User(BaseModel):
    model_config = ConfigDict(extra="ignore")
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    email: EmailStr
    name: str
    picture: Optional[str] = None
    user_type: UserType = UserType.REPRESENTANTE
    board_position: Optional[BoardPosition] = None
    vocalia_ids: List[str] = Field(default_factory=list)
    is_active: bool = True
    university_id: Optional[str] = None
    created_at: datetime = Field(default_factory=lambda: datetime.now(timezone.utc))


class University(BaseModel):
    model_config = ConfigDict(extra="ignore")
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    name: str
    is_active: bool = True
    zone: Optional[str] = None
    acronym: Optional[str] = None
    region: Optional[str] = None
    website_url: Optional[str] = None
    center_name: Optional[str] = None
    center_url: Optional[str] = None
    is_ritsi_member: Optional[bool] = None
    source_key: Optional[str] = None
    created_at: datetime = Field(default_factory=lambda: datetime.now(timezone.utc))


class Vocalia(BaseModel):
    model_config = ConfigDict(extra="ignore")
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    name: str
    description: Optional[str] = None
    board_member_id: str
    created_at: datetime = Field(default_factory=lambda: datetime.now(timezone.utc))


class Category(BaseModel):
    model_config = ConfigDict(extra="ignore")
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    name: str
    created_at: datetime = Field(default_factory=lambda: datetime.now(timezone.utc))


class ContentFile(BaseModel):
    model_config = ConfigDict(extra="ignore")
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    file_type: FileType = FileType.LINK
    url: str
    title: str
    description: Optional[str] = None


class Question(BaseModel):
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    question_text: str = Field(min_length=2, max_length=1000)
    question_type: QuestionType
    options: List[str] = Field(default_factory=list)
    correct_answers: List[int] = Field(default_factory=list)

    @model_validator(mode="after")
    def validate_answers(self):
        if len(self.options) < 2 or any(not option.strip() for option in self.options):
            raise ValueError("Cada pregunta necesita al menos dos opciones no vacias")
        if not self.correct_answers or len(set(self.correct_answers)) != len(self.correct_answers):
            raise ValueError("Configura al menos una respuesta correcta sin duplicados")
        if any(index < 0 or index >= len(self.options) for index in self.correct_answers):
            raise ValueError("Una respuesta correcta no corresponde a ninguna opcion")
        if self.question_type in {QuestionType.MULTIPLE_CHOICE, QuestionType.TRUE_FALSE} and len(self.correct_answers) != 1:
            raise ValueError("Este tipo de pregunta admite una unica respuesta correcta")
        return self


class Quiz(BaseModel):
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    title: str
    questions: List[Question] = Field(min_length=1)
    passing_percentage: float = Field(default=70.0, ge=0, le=100)


class TrainingContent(BaseModel):
    model_config = ConfigDict(extra="ignore")
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    external_id: Optional[str] = None
    source_document_id: Optional[str] = None
    source_code: Optional[str] = None
    source_sheet: Optional[str] = None
    academic_year: Optional[str] = None
    title: str
    description: Optional[str] = None
    training_date: Optional[str] = None
    audience: Optional[str] = None
    duration_minutes: Optional[int] = None
    attendees: Optional[int] = None
    rating: Optional[float] = None
    trainer_names: List[str] = Field(default_factory=list)
    tags: List[str] = Field(default_factory=list)
    status: ContentStatus = ContentStatus.PENDING
    is_public: bool = False
    files: List[ContentFile] = Field(default_factory=list)
    category_ids: List[str] = Field(default_factory=list)
    quizzes: List[Quiz] = Field(default_factory=list)
    created_by: str
    created_at: datetime = Field(default_factory=lambda: datetime.now(timezone.utc))
    updated_at: Optional[datetime] = None
    review_note: Optional[str] = None
    reviewed_by: Optional[str] = None
    reviewed_at: Optional[datetime] = None


class QuestionView(BaseModel):
    """Question contract safe to expose before an attempt is submitted."""

    id: str
    question_text: str
    question_type: QuestionType
    options: List[str] = Field(default_factory=list)


class QuizView(BaseModel):
    id: str
    title: str
    questions: List[QuestionView] = Field(default_factory=list)
    passing_percentage: float = 70.0


class TrainingContentView(TrainingContent):
    """Learner-safe catalogue contract; answer keys are deliberately absent."""

    quizzes: List[QuizView] = Field(default_factory=list)


class ContentAssignment(BaseModel):
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    content_id: str
    assigned_to_user_ids: List[str] = Field(default_factory=list)
    assigned_to_all_representatives: bool = False
    assigned_by: str
    created_at: datetime = Field(default_factory=lambda: datetime.now(timezone.utc))


class Enrollment(BaseModel):
    model_config = ConfigDict(extra="ignore")
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    content_id: str
    user_id: str
    assigned_by: str
    assignment_ids: List[str] = Field(default_factory=list)
    source: EnrollmentSource = EnrollmentSource.ASSIGNMENT
    status: EnrollmentStatus = EnrollmentStatus.ASSIGNED
    created_at: datetime = Field(default_factory=lambda: datetime.now(timezone.utc))
    started_at: Optional[datetime] = None
    completed_at: Optional[datetime] = None
    updated_at: datetime = Field(default_factory=lambda: datetime.now(timezone.utc))


class SynchronousSession(BaseModel):
    model_config = ConfigDict(extra="ignore")
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    content_id: str
    title: str
    description: Optional[str] = None
    starts_at: datetime
    ends_at: datetime
    timezone: str = DEFAULT_TIMEZONE
    status: SynchronousSessionStatus = SynchronousSessionStatus.SCHEDULED
    participant_user_ids: List[str] = Field(default_factory=list)
    join_window_minutes: int = DEFAULT_JOIN_WINDOW_MINUTES
    created_by: str
    created_at: datetime = Field(default_factory=lambda: datetime.now(timezone.utc))
    updated_at: datetime = Field(default_factory=lambda: datetime.now(timezone.utc))


class SynchronousSessionView(BaseModel):
    model_config = ConfigDict(extra="ignore")
    id: str
    content_id: str
    title: str
    description: Optional[str] = None
    starts_at: datetime
    ends_at: datetime
    timezone: str
    status: SynchronousSessionStatus
    participant_count: int = 0
    join_window_minutes: int
    created_by: str
    created_at: datetime
    updated_at: datetime
    can_join: bool = False
    join_unavailable_reason: Optional[str] = None
    join_available_from: datetime
    can_manage: bool = False


class ActivityLog(BaseModel):
    id: str = Field(default_factory=lambda: str(uuid.uuid4()))
    timestamp: datetime = Field(default_factory=lambda: datetime.now(timezone.utc))
    actor_id: str
    actor_name: str
    action: str
    target_user_id: str
    target_user_name: str
    details: Optional[Dict[str, Any]] = None


class LoginRequest(BaseModel):
    email: EmailStr
    password: str = Field(min_length=8, max_length=128)


class RegisterRequest(LoginRequest):
    name: str = Field(min_length=2, max_length=120)
    university_id: str


class UserCreate(LoginRequest):
    name: str = Field(min_length=2, max_length=120)
    user_type: UserType = UserType.REPRESENTANTE
    university_id: Optional[str] = None


class UpdateUserRequest(BaseModel):
    name: Optional[str] = Field(default=None, min_length=2, max_length=120)
    user_type: Optional[UserType] = None
    university_id: Optional[str] = None
    is_active: Optional[bool] = None


class BoardMemberInput(BaseModel):
    user_id: str
    position: BoardPosition


class BoardUpdateRequest(BaseModel):
    members: List[BoardMemberInput]


class VocaliaCreate(BaseModel):
    name: str = Field(min_length=2, max_length=120)
    description: Optional[str] = Field(default=None, max_length=600)
    board_member_id: str


class AssignUsersRequest(BaseModel):
    user_ids: List[str]


class UniversityCreate(BaseModel):
    name: str = Field(min_length=2, max_length=160)
    zone: Optional[str] = None


class UpdateStatusRequest(BaseModel):
    is_active: bool


class CategoryCreate(BaseModel):
    name: str = Field(min_length=2, max_length=100)


class ContentFileCreate(BaseModel):
    file_type: FileType = FileType.LINK
    url: str
    title: str
    description: Optional[str] = None

    @field_validator("url")
    @classmethod
    def validate_url(cls, value: str) -> str:
        if not value.startswith(("https://", "http://")):
            raise ValueError("La URL debe comenzar por http:// o https://")
        return value


class TrainingContentCreate(BaseModel):
    title: str = Field(min_length=2, max_length=240)
    description: Optional[str] = None
    is_public: bool = False
    category_ids: List[str] = Field(default_factory=list)
    files: List[ContentFileCreate] = Field(default_factory=list)
    quizzes: List[Quiz] = Field(default_factory=list)
    training_date: Optional[str] = None
    audience: Optional[str] = None
    duration_minutes: Optional[int] = None
    trainer_names: List[str] = Field(default_factory=list)
    tags: List[str] = Field(default_factory=list)


class AssignContentRequest(BaseModel):
    content_id: str
    user_ids: List[str] = Field(default_factory=list)
    assign_to_all_representatives: bool = False


class SessionCreateRequest(BaseModel):
    content_id: str
    title: str = Field(min_length=2, max_length=240)
    description: Optional[str] = Field(default=None, max_length=2000)
    starts_at: datetime
    ends_at: datetime
    timezone: str = DEFAULT_TIMEZONE
    participant_user_ids: List[str] = Field(default_factory=list)
    join_window_minutes: int = Field(default=DEFAULT_JOIN_WINDOW_MINUTES, ge=0, le=180)
    meeting_url: Optional[str] = None

    @field_validator("meeting_url")
    @classmethod
    def validate_meeting_url(cls, value: Optional[str]) -> Optional[str]:
        if value is not None and not value.startswith("https://"):
            raise ValueError("La URL de la videollamada debe comenzar por https://")
        return value


class SessionUpdateRequest(BaseModel):
    title: Optional[str] = Field(default=None, min_length=2, max_length=240)
    description: Optional[str] = Field(default=None, max_length=2000)
    starts_at: Optional[datetime] = None
    ends_at: Optional[datetime] = None
    timezone: Optional[str] = None
    status: Optional[SynchronousSessionStatus] = None
    join_window_minutes: Optional[int] = Field(default=None, ge=0, le=180)
    meeting_url: Optional[str] = None

    @field_validator("meeting_url")
    @classmethod
    def validate_meeting_url(cls, value: Optional[str]) -> Optional[str]:
        if value is not None and not value.startswith("https://"):
            raise ValueError("La URL de la videollamada debe comenzar por https://")
        return value


class SessionParticipantsRequest(BaseModel):
    user_ids: List[str] = Field(default_factory=list)


class ImportFormationsRequest(BaseModel):
    spreadsheet_id: str = SHEET_ID
    publish: bool = True
    dry_run: bool = False


class NotificationPreferencesRequest(BaseModel):
    in_app_enabled: bool = True
    telegram_enabled: bool = True
    assignment_notifications: bool = True
    progress_notifications: bool = True
    session_notifications: bool = True


class ResetPasswordRequest(BaseModel):
    token: str = Field(min_length=20, max_length=300)
    password: str = Field(min_length=8, max_length=128)


class ProfileUpdateRequest(BaseModel):
    name: str = Field(min_length=2, max_length=120)
    picture: Optional[str] = Field(default=None, max_length=1000)

    @field_validator("picture")
    @classmethod
    def validate_picture(cls, value: Optional[str]) -> Optional[str]:
        if value and not value.startswith("https://"):
            raise ValueError("La imagen de perfil debe usar HTTPS")
        return value


class ChangePasswordRequest(BaseModel):
    current_password: str = Field(min_length=8, max_length=128)
    new_password: str = Field(min_length=8, max_length=128)


class AttendanceEntry(BaseModel):
    user_id: str
    status: str = Field(pattern="^(absent|partial|attended|excused)$")
    minutes: int = Field(default=0, ge=0, le=1440)
    notes: Optional[str] = Field(default=None, max_length=500)


class AttendanceRequest(BaseModel):
    entries: List[AttendanceEntry]


class ContentUpdateRequest(BaseModel):
    title: Optional[str] = Field(default=None, min_length=2, max_length=240)
    description: Optional[str] = None
    is_public: Optional[bool] = None
    category_ids: Optional[List[str]] = None
    files: Optional[List[ContentFileCreate]] = None
    quizzes: Optional[List[Quiz]] = None
    training_date: Optional[str] = None
    audience: Optional[str] = None
    duration_minutes: Optional[int] = None
    trainer_names: Optional[List[str]] = None
    tags: Optional[List[str]] = None


class ReviewRequest(BaseModel):
    action: str = Field(pattern="^(publish|request_changes|archive)$")
    note: Optional[str] = Field(default=None, max_length=1000)


class LearningPathRequest(BaseModel):
    title: str = Field(min_length=2, max_length=240)
    description: Optional[str] = Field(default=None, max_length=2000)
    content_ids: List[str] = Field(min_length=1, max_length=100)
    is_active: bool = True


class LearningPathAssignmentRequest(BaseModel):
    user_ids: List[str] = Field(min_length=1, max_length=5000)


class AccountDeletionRequest(BaseModel):
    password: str = Field(min_length=8, max_length=128)


class MarkFileCompletedRequest(BaseModel):
    content_id: str
    file_id: str


class SubmitQuizRequest(BaseModel):
    content_id: str
    quiz_id: str
    answers: Dict[str, List[int]]


def serialize(model: BaseModel) -> Dict[str, Any]:
    return model.model_dump(mode="json")


def mongo_serialize(model: BaseModel) -> Dict[str, Any]:
    """Keep native datetimes for new MongoDB records."""

    return model.model_dump(mode="python")


def token_digest(token: str) -> str:
    return hashlib.sha256(token.encode("utf-8")).hexdigest()


async def enqueue_event(event_type: str, aggregate_type: str, aggregate_id: str, payload: Dict[str, Any], dedupe_key: str) -> None:
    now = datetime.now(timezone.utc)
    await db.outbox_events.update_one(
        {"dedupe_key": dedupe_key},
        {"$setOnInsert": {
            "id": str(uuid.uuid4()), "event_type": event_type,
            "aggregate_type": aggregate_type, "aggregate_id": aggregate_id,
            "payload": payload, "dedupe_key": dedupe_key, "status": "pending",
            "attempts": 0, "occurred_at": now, "available_at": now,
        }}, upsert=True,
    )


async def enforce_login_rate_limit(email: str, client_host: str) -> str:
    """Fixed-window limiter stored in Mongo so all API replicas share state."""
    now = datetime.now(timezone.utc)
    window = now.replace(second=0, microsecond=0)
    key = token_digest(f"{client_host.lower()}:{email.lower()}:{window.isoformat()}")
    attempt = await db.auth_rate_limits.find_one_and_update(
        {"key": key},
        {"$inc": {"attempts": 1}, "$setOnInsert": {"key": key, "created_at": now, "expires_at": now + timedelta(minutes=15)}},
        upsert=True, return_document=ReturnDocument.AFTER,
    )
    if attempt.get("attempts", 0) > 10:
        raise HTTPException(status_code=429, detail="Demasiados intentos. Espera antes de volver a probar", headers={"Retry-After": "60"})
    return key


def browser_mutation_origin_is_trusted(origin: Optional[str], fetch_site: Optional[str]) -> bool:
    if origin:
        return origin in ALLOWED_ORIGINS
    return (fetch_site or "").lower() != "cross-site"


def parse_utc_datetime(value: Any) -> datetime:
    if isinstance(value, datetime):
        parsed = value
    else:
        parsed = datetime.fromisoformat(str(value).replace("Z", "+00:00"))
    if parsed.tzinfo is None or parsed.utcoffset() is None:
        raise ValueError("La fecha debe incluir zona horaria")
    return parsed.astimezone(timezone.utc)


def parse_stored_utc_datetime(value: Any) -> datetime:
    """Normalize persisted BSON datetimes, which MongoDB returns as naive UTC."""
    if isinstance(value, datetime) and (value.tzinfo is None or value.utcoffset() is None):
        return value.replace(tzinfo=timezone.utc)
    return parse_utc_datetime(value)


def validate_session_schedule(starts_at: datetime, ends_at: datetime, timezone_name: str) -> tuple[datetime, datetime]:
    try:
        ZoneInfo(timezone_name)
    except ZoneInfoNotFoundError as error:
        raise ValueError("Zona horaria IANA no valida") from error
    start_utc = parse_utc_datetime(starts_at)
    end_utc = parse_utc_datetime(ends_at)
    if end_utc <= start_utc:
        raise ValueError("La sesion debe terminar despues de comenzar")
    return start_utc, end_utc


def calculate_progress_completion(content: Dict[str, Any], progress: Dict[str, Any]) -> bool:
    """A content item completes when every resource is opened and every quiz is passed."""

    required_files = {item.get("id") for item in content.get("files", []) if item.get("id")}
    required_quizzes = {item.get("id") for item in content.get("quizzes", []) if item.get("id")}
    if not required_files and not required_quizzes:
        return False
    completed_files = set(progress.get("files_completed", []))
    quiz_results = progress.get("quizzes_completed", {})
    return required_files.issubset(completed_files) and all(
        bool(quiz_results.get(quiz_id, {}).get("passed")) for quiz_id in required_quizzes
    )


def session_join_state(session: Dict[str, Any], now: Optional[datetime] = None) -> tuple[bool, Optional[str]]:
    """Pure time/state gate. Identity and enrolment are checked separately."""

    current = parse_stored_utc_datetime(now or datetime.now(timezone.utc))
    starts_at = parse_stored_utc_datetime(session["starts_at"])
    ends_at = parse_stored_utc_datetime(session["ends_at"])
    status = SynchronousSessionStatus(session.get("status", SynchronousSessionStatus.SCHEDULED.value))
    if status == SynchronousSessionStatus.CANCELLED:
        return False, "La sesion esta cancelada"
    if status == SynchronousSessionStatus.DRAFT:
        return False, "La sesion aun no esta programada"
    if status == SynchronousSessionStatus.ENDED or current >= ends_at:
        return False, "La sesion ha finalizado"
    available_from = starts_at - timedelta(minutes=int(session.get("join_window_minutes", DEFAULT_JOIN_WINDOW_MINUTES)))
    if current < available_from:
        return False, "El acceso todavia no esta disponible"
    return True, None


def meeting_url_storage_fields(meeting_url: str) -> Dict[str, str]:
    if not MEETING_URL_ENCRYPTION_KEY:
        raise RuntimeError("MEETING_URL_ENCRYPTION_KEY no esta configurada")
    key = base64.urlsafe_b64encode(hashlib.sha256(MEETING_URL_ENCRYPTION_KEY.encode("utf-8")).digest())
    ciphertext = Fernet(key).encrypt(meeting_url.encode("utf-8")).decode("ascii")
    return {
        "meeting_url_ciphertext": ciphertext,
        "meeting_url_hash": token_digest(meeting_url),
    }


def read_meeting_url(session: Dict[str, Any]) -> Optional[str]:
    ciphertext = session.get("meeting_url_ciphertext")
    if ciphertext:
        if not MEETING_URL_ENCRYPTION_KEY:
            return None
        key = base64.urlsafe_b64encode(hashlib.sha256(MEETING_URL_ENCRYPTION_KEY.encode("utf-8")).digest())
        try:
            return Fernet(key).decrypt(ciphertext.encode("ascii")).decode("utf-8")
        except (InvalidToken, ValueError):
            return None
    # Transitional read support. initialize_database encrypts and removes this
    # field whenever an encryption key is configured.
    return session.get("meeting_url")


def password_hash(password: str) -> str:
    return bcrypt.hashpw(password.encode(), bcrypt.gensalt()).decode()


def password_matches(password: str, hashed: str) -> bool:
    try:
        return bcrypt.checkpw(password.encode(), hashed.encode())
    except (ValueError, AttributeError):
        return False


def normalize(value: Any) -> str:
    text = "" if value is None else str(value).strip().lower()
    return "".join(c for c in unicodedata.normalize("NFD", text) if unicodedata.category(c) != "Mn")


def modernize_text(value: Any) -> str:
    text = str(value or "").strip()
    replacements = (
        (r"(?i)coordinaciones\s+tem[aá]ticas", "Vocalías"),
        (r"(?i)coordinaci[oó]n\s+tem[aá]tica", "Vocalía"),
        (r"(?i)comisiones\s+tem[aá]ticas", "Vocalías"),
        (r"(?i)comisi[oó]n\s+tem[aá]tica", "Vocalía"),
        (r"(?i)coordinador(?:a)?\s+de\s+la\s+escuela\s+de\s+formaci[oó]n", "Vocalía de Formación"),
        (r"(?i)coordinaci[oó]n\s+de\s+la\s+escuela\s+de\s+formaci[oó]n", "Vocalía de Formación"),
        (r"(?i)escuela\s+de\s+formaci[oó]n", "Vocalía de Formación"),
    )
    for pattern, replacement in replacements:
        text = re.sub(pattern, replacement, text)
    return "Vocalías" if normalize(text) == "ct" else text


RITSI_REGIONS = {
    1: ("Andalucía", "I"), 2: ("Aragón", "IV"), 3: ("Asturias", "III"),
    4: ("Islas Baleares", "V"), 5: ("País Vasco", "III"), 6: ("Islas Canarias", "IV"),
    7: ("Cantabria", "III"), 8: ("Castilla y León", "IV"), 9: ("Castilla-La Mancha", "I"),
    10: ("Cataluña", "V"), 11: ("Extremadura", "IV"), 12: ("Galicia", "III"),
    13: ("La Rioja", "III"), 14: ("Madrid", "II"), 15: ("Murcia", "I"),
    16: ("Navarra", "III"), 17: ("Valencia", "V"),
}


def clean_html(value: str) -> str:
    return re.sub(r"\s+", " ", html_lib.unescape(re.sub(r"<[^>]+>", " ", value))).strip()


def parse_ritsi_universities(region_id: int, markup: str) -> List[Dict[str, Any]]:
    region, zone = RITSI_REGIONS[region_id]
    cards = re.findall(
        r'<div class="mdc-card (no-member|member)">(.*?)(?=<div class="item">|\Z)',
        markup,
        flags=re.IGNORECASE | re.DOTALL,
    )
    universities = []
    for membership, card in cards:
        name_match = re.search(r'class="university-name">.*?<b>(.*?)</b>', card, flags=re.IGNORECASE | re.DOTALL)
        if not name_match:
            continue
        name = clean_html(name_match.group(1))
        acronym_match = re.search(r"\(([A-ZÀ-Ü0-9-]{2,12})\)\s*$", name)
        website_match = re.search(r'class="mdc-card__media[^>]*href="([^"]+)"', card, flags=re.IGNORECASE | re.DOTALL)
        center_match = re.search(r'class="center-name">.*?<a href="([^"]+)"[^>]*>.*?<u>(.*?)</u>', card, flags=re.IGNORECASE | re.DOTALL)
        universities.append({
            "name": name,
            "acronym": acronym_match.group(1) if acronym_match else None,
            "region": region,
            "zone": zone,
            "website_url": html_lib.unescape(website_match.group(1)) if website_match else None,
            "center_url": html_lib.unescape(center_match.group(1)) if center_match else None,
            "center_name": clean_html(center_match.group(2)) if center_match else None,
            "is_ritsi_member": membership.lower() == "member",
            "is_active": membership.lower() == "member",
            "source_key": f"ritsi:{normalize(name)}",
        })
    return universities


async def import_ritsi_universities() -> Dict[str, int]:
    async def fetch_region(region_id: int):
        url = f"https://ritsi.org/index.php?freemap_get_state_info={region_id}&map_id=0"
        response = await asyncio.to_thread(requests.post, url, timeout=30)
        response.raise_for_status()
        # The endpoint serves UTF-8 HTML without declaring its charset, so
        # requests otherwise falls back to ISO-8859-1 and corrupts accents.
        return parse_ritsi_universities(region_id, response.content.decode("utf-8"))

    region_results = await asyncio.gather(*(fetch_region(region_id) for region_id in RITSI_REGIONS))
    unique: Dict[str, Dict[str, Any]] = {}
    for item in (university for region in region_results for university in region):
        key = item["source_key"]
        if key not in unique or item["is_ritsi_member"]:
            unique[key] = item
    created = updated = 0
    for item in unique.values():
        matches = [
            {"source_key": item["source_key"]},
            {"name": {"$regex": f"^{re.escape(item['name'])}$", "$options": "i"}},
        ]
        # The URL also lets a re-import repair records created before the
        # source encoding was handled explicitly.
        if item.get("website_url"):
            matches.append({"website_url": item["website_url"]})
        existing = await db.universities.find_one({"$or": matches})
        if existing:
            await db.universities.update_one({"_id": existing["_id"]}, {"$set": item})
            updated += 1
        else:
            university = University(**item)
            await db.universities.insert_one(serialize(university))
            created += 1
    return {"found": len(unique), "created": created, "updated": updated}


def validate_board_members(members: List[BoardMemberInput]) -> None:
    if not 5 <= len(members) <= 7:
        raise ValueError("La Junta Directiva debe tener entre 5 y 7 miembros")
    user_ids = [member.user_id for member in members]
    positions = [member.position for member in members]
    if len(set(user_ids)) != len(user_ids):
        raise ValueError("Una persona no puede ocupar dos cargos de Junta Directiva")
    if len(set(positions)) != len(positions):
        raise ValueError("Cada cargo de Junta Directiva solo puede asignarse una vez")
    missing = REQUIRED_BOARD_POSITIONS - set(positions)
    if missing:
        raise ValueError("Faltan cargos obligatorios: " + ", ".join(sorted(item.value for item in missing)))


def infer_file_type(title: str, url: str) -> FileType:
    value = normalize(f"{title} {url}")
    if "grabacion" in value or "youtube" in value or "youtu.be" in value or "video" in value:
        return FileType.VIDEO
    if ".pdf" in value or "informe" in value:
        return FileType.PDF
    if "presentation" in value or "presentacion" in value or "canva" in value:
        return FileType.PRESENTATION
    if any(ext in value for ext in (".png", ".jpg", ".jpeg", ".webp")):
        return FileType.IMAGE
    if "document" in value or "documento" in value:
        return FileType.DOCUMENT
    return FileType.LINK


def cell_url(cell: Any) -> Optional[str]:
    if cell.hyperlink:
        return cell.hyperlink.target
    match = re.search(r'HYPERLINK\("([^"]+)', str(cell.value or ""), flags=re.IGNORECASE)
    return match.group(1) if match else None


def clean_int(value: Any) -> Optional[int]:
    try:
        return int(float(str(value).replace(",", ".")))
    except (TypeError, ValueError):
        return None


def clean_float(value: Any) -> Optional[float]:
    try:
        return float(str(value).replace(",", "."))
    except (TypeError, ValueError):
        return None


def clean_date(value: Any) -> Optional[str]:
    if isinstance(value, (datetime, date)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    text = str(value or "").strip()
    if not text or text == "-":
        return None
    for pattern in ("%d/%m/%Y", "%d/%m/%y"):
        try:
            return datetime.strptime(text, pattern).date().isoformat()
        except ValueError:
            pass
    return text


def stable_import_resource_id(content_external_id: str, url: str, occurrence: int = 0) -> str:
    key = f"ritsi-training:{content_external_id}:resource:{url}:{occurrence}"
    return str(uuid.uuid5(uuid.NAMESPACE_URL, key))


IMPORTED_SOURCE_FIELDS = {
    "source_document_id",
    "source_code",
    "source_sheet",
    "academic_year",
    "title",
    "training_date",
    "audience",
    "duration_minutes",
    "attendees",
    "rating",
    "trainer_names",
    "tags",
    "files",
}


def imported_source_update(content: TrainingContent) -> Dict[str, Any]:
    payload = serialize(content)
    return {field: payload[field] for field in IMPORTED_SOURCE_FIELDS}


def parse_sheet_workbook(binary: bytes, created_by: str, publish: bool) -> List[TrainingContent]:
    workbook = load_workbook(io.BytesIO(binary), data_only=False, read_only=False)
    results: List[TrainingContent] = []
    for worksheet in workbook.worksheets:
        if normalize(worksheet.title) in {"hoja 5", "siglas"}:
            continue
        header_row, headers = None, []
        for row_index in range(1, min(worksheet.max_row, 8) + 1):
            values = [normalize(worksheet.cell(row_index, column).value) for column in range(1, worksheet.max_column + 1)]
            if any("codigo" in value and "formacion" in value for value in values):
                header_row, headers = row_index, values
                break
        if not header_row:
            continue

        def index_of(predicate, default=None):
            return next((index for index, value in enumerate(headers) if predicate(value)), default)

        code_i = index_of(lambda value: "codigo" in value)
        date_i = index_of(lambda value: value == "fecha")
        audience_i = index_of(lambda value: "dirigido" in value)
        attendees_i = index_of(lambda value: "asistentes" in value)
        duration_i = index_of(lambda value: "duracion" in value)
        rating_i = index_of(lambda value: "valoracion" in value)
        files_i = index_of(lambda value: "archivos adjuntos" in value)
        tags_i = index_of(lambda value: "etiquetas" in value)
        name_indices = [i for i, value in enumerate(headers) if value == "nombre"]
        title_i = next((i for i in name_indices if duration_i is not None and i > duration_i), name_indices[0] if name_indices else None)
        trainer_indices = [i for i in name_indices if title_i is not None and i > title_i and (files_i is None or i < files_i)]
        if code_i is None or title_i is None:
            continue
        empty_streak = 0
        for row_index in range(header_row + 1, worksheet.max_row + 1):
            cells = list(worksheet[row_index])
            code = str(cells[code_i].value or "").strip()
            title = modernize_text(cells[title_i].value)
            if not code and not title:
                empty_streak += 1
                if empty_streak >= 20:
                    break
                continue
            empty_streak = 0
            if normalize(code) in {"total", "totales"} or normalize(title) in {"total", "totales"}:
                continue
            content_external_id = f"{worksheet.title}:{code or title}"
            files: List[ContentFile] = []
            resource_occurrences: Dict[str, int] = {}
            if files_i is not None:
                end_i = tags_i if tags_i is not None else min(len(cells), files_i + 6)
                for cell in cells[files_i:end_i]:
                    url = cell_url(cell)
                    if url:
                        file_title = str(cell.value or "Recurso").strip()
                        occurrence = resource_occurrences.get(url, 0)
                        resource_occurrences[url] = occurrence + 1
                        files.append(ContentFile(
                            id=stable_import_resource_id(content_external_id, url, occurrence),
                            file_type=infer_file_type(file_title, url),
                            url=url,
                            title=file_title,
                        ))
            trainers = []
            for index in trainer_indices:
                value = str(cells[index].value or "").strip()
                if value and value not in trainers:
                    trainers.append(value)
            raw_tags = modernize_text(cells[tags_i].value) if tags_i is not None else ""
            tags = [modernize_text(item) for item in re.split(r"[,;\n]", raw_tags) if item.strip()]
            results.append(TrainingContent(
                external_id=content_external_id, source_code=code or None,
                source_sheet=worksheet.title, academic_year=worksheet.title, title=title or code,
                description=raw_tags.strip() or None,
                training_date=clean_date(cells[date_i].value) if date_i is not None else None,
                audience=(modernize_text(cells[audience_i].value) or None) if audience_i is not None else None,
                duration_minutes=clean_int(cells[duration_i].value) if duration_i is not None else None,
                attendees=clean_int(cells[attendees_i].value) if attendees_i is not None else None,
                rating=clean_float(cells[rating_i].value) if rating_i is not None else None,
                trainer_names=trainers, tags=tags, files=files,
                status=ContentStatus.PUBLISHED if publish else ContentStatus.PENDING,
                is_public=True, created_by=created_by,
            ))
    return results


async def import_formations(spreadsheet_id: str, created_by: str, publish: bool = True, dry_run: bool = False) -> Dict[str, int]:
    response = await asyncio.to_thread(requests.get, f"https://docs.google.com/spreadsheets/d/{spreadsheet_id}/export?format=xlsx", timeout=60)
    response.raise_for_status()
    contents = await asyncio.to_thread(parse_sheet_workbook, response.content, created_by, publish)
    created = updated = unchanged = 0
    for content in contents:
        legacy_external_id = content.external_id
        scoped_external_id = f"{spreadsheet_id}:{legacy_external_id}"
        existing = await db.training_contents.find_one({"$or": [
            {"external_id": scoped_external_id},
            {"external_id": legacy_external_id, "source_document_id": spreadsheet_id},
            {"external_id": legacy_external_id, "source_document_id": {"$exists": False}},
        ]}, {"id": 1, **{field: 1 for field in IMPORTED_SOURCE_FIELDS}})
        content.source_document_id = spreadsheet_id
        if existing:
            source_update = imported_source_update(content)
            changed = any(existing.get(field) != value for field, value in source_update.items())
            if changed:
                if not dry_run:
                    await db.training_contents.update_one({"id": existing["id"]}, {"$set": source_update})
                updated += 1
            else:
                unchanged += 1
        else:
            content.external_id = scoped_external_id
            if not dry_run:
                await db.training_contents.insert_one(serialize(content))
            created += 1
    return {"found": len(contents), "created": created, "updated": updated, "unchanged": unchanged, "dry_run": dry_run}


async def log_activity(actor: User, action: str, target: Dict[str, Any], details: Optional[Dict[str, Any]] = None):
    entry = ActivityLog(actor_id=actor.id, actor_name=actor.name, action=action,
                        target_user_id=str(target.get("id", "system")),
                        target_user_name=str(target.get("name", target.get("title", "Sistema"))), details=details)
    await db.activity_logs.insert_one(serialize(entry))


async def get_current_user(session_token: Optional[str] = Cookie(default=None)) -> User:
    if not session_token:
        raise HTTPException(status_code=401, detail="No autorizado")
    session = await db.user_sessions.find_one({"$or": [
        {"session_token_hash": token_digest(session_token)},
        {"session_token": session_token},  # Legacy records during the migration window.
    ]})
    if not session:
        raise HTTPException(status_code=401, detail="Sesión inválida o expirada")
    try:
        expired = parse_stored_utc_datetime(session["expires_at"]) <= datetime.now(timezone.utc)
    except (KeyError, TypeError, ValueError):
        expired = True
    if expired:
        await db.user_sessions.delete_one({"_id": session["_id"]})
        raise HTTPException(status_code=401, detail="Sesión inválida o expirada")
    user = await db.users.find_one({"id": session["user_id"]}, {"_id": 0, "password_hash": 0})
    if not user:
        await db.user_sessions.delete_one({"_id": session["_id"]})
        raise HTTPException(status_code=401, detail="Sesion invalida")
    if not user.get("is_active", True):
        await db.user_sessions.delete_many({"user_id": session["user_id"]})
        raise HTTPException(status_code=403, detail="La cuenta está desactivada")
    return User(**user)


async def create_session(user: User) -> str:
    token = secrets.token_urlsafe(48)
    now = datetime.now(timezone.utc)
    await db.user_sessions.insert_one({"id": str(uuid.uuid4()), "user_id": user.id,
        "session_token_hash": token_digest(token), "expires_at": now + timedelta(days=SESSION_DAYS),
        "created_at": now})
    return token


def session_response(user: User, token: str, code: int = 200) -> JSONResponse:
    response = JSONResponse(serialize(user), status_code=code)
    response.set_cookie("session_token", token, httponly=True, secure=COOKIE_SECURE, samesite="lax", max_age=SESSION_DAYS * 86400, path="/")
    return response


@api_router.get("/health")
async def health():
    try:
        await db.command("ping")
    except Exception as error:
        logger.error("MongoDB no esta disponible en el healthcheck: %s", error)
        raise HTTPException(status_code=503, detail="Base de datos no disponible") from error
    return {"status": "ok", "database": "ok"}


@api_router.post("/auth/register")
async def register(request: RegisterRequest):
    if await db.users.find_one({"email": request.email.lower()}):
        raise HTTPException(status_code=409, detail="Ya existe una cuenta con ese correo")
    if not await db.universities.find_one({"id": request.university_id, "is_active": True}):
        raise HTTPException(status_code=400, detail="Selecciona una universidad activa")
    user = User(email=request.email.lower(), name=request.name.strip(), university_id=request.university_id)
    payload = serialize(user)
    payload["password_hash"] = password_hash(request.password)
    await db.users.insert_one(payload)
    return session_response(user, await create_session(user), 201)


@api_router.post("/auth/login")
async def login(request: LoginRequest, http_request: Request):
    rate_key = await enforce_login_rate_limit(request.email, http_request.client.host if http_request.client else "unknown")
    stored = await db.users.find_one({"email": request.email.lower()})
    if not stored or not password_matches(request.password, stored.get("password_hash", "")):
        raise HTTPException(status_code=401, detail="Correo o contraseña incorrectos")
    if not stored.get("is_active", True):
        raise HTTPException(status_code=403, detail="La cuenta está desactivada")
    user = User(**stored)
    await db.auth_rate_limits.delete_one({"key": rate_key})
    return session_response(user, await create_session(user))


@api_router.get("/auth/me", response_model=User)
async def get_me(current_user: User = Depends(get_current_user)):
    return current_user


@api_router.patch("/profile", response_model=User)
async def update_profile(request: ProfileUpdateRequest, current_user: User = Depends(get_current_user)):
    changes = {"name": request.name.strip(), "picture": request.picture, "updated_at": datetime.now(timezone.utc)}
    await db.users.update_one({"id": current_user.id}, {"$set": changes})
    return User(**await db.users.find_one({"id": current_user.id}, {"_id": 0, "password_hash": 0}))


@api_router.post("/profile/change-password")
async def change_own_password(request: ChangePasswordRequest, current_user: User = Depends(get_current_user)):
    stored = await db.users.find_one({"id": current_user.id})
    if not stored or not password_matches(request.current_password, stored.get("password_hash", "")):
        raise HTTPException(status_code=400, detail="La contraseña actual no es correcta")
    await db.users.update_one({"id": current_user.id}, {"$set": {"password_hash": password_hash(request.new_password)}})
    await db.user_sessions.delete_many({"user_id": current_user.id})
    response = JSONResponse({"success": True, "reauthentication_required": True})
    response.delete_cookie("session_token", path="/")
    return response


@api_router.get("/privacy/export")
async def export_personal_data(current_user: User = Depends(get_current_user)):
    async def records(collection, query, projection=None):
        return await collection.find(query, projection or {"_id": 0}).to_list(10000)
    return {
        "exported_at": datetime.now(timezone.utc),
        "profile": serialize(current_user),
        "enrollments": await records(db.enrollments, {"user_id": current_user.id}),
        "progress": await records(db.user_progress, {"user_id": current_user.id}),
        "attendance": await records(db.session_attendance, {"user_id": current_user.id}),
        "certificates": await records(db.certificates, {"user_id": current_user.id}),
        "notification_preferences": await db.notification_preferences.find_one({"user_id": current_user.id}, {"_id": 0}) or {},
        "telegram_linked": bool(await db.telegram_bindings.find_one({"user_id": current_user.id, "active": True}, {"id": 1})),
    }


@api_router.post("/privacy/deletion-request")
async def request_account_deletion(request: AccountDeletionRequest, current_user: User = Depends(get_current_user)):
    stored = await db.users.find_one({"id": current_user.id})
    if not stored or not password_matches(request.password, stored.get("password_hash", "")):
        raise HTTPException(status_code=400, detail="La contraseña no es correcta")
    if current_user.user_type in {UserType.ADMIN, UserType.JUNTA_DIRECTIVA}:
        raise HTTPException(status_code=409, detail="Transfiere primero las responsabilidades de gobierno de la cuenta")
    now = datetime.now(timezone.utc)
    await db.privacy_requests.update_one(
        {"user_id": current_user.id, "status": "pending"},
        {"$setOnInsert": {"id": str(uuid.uuid4()), "user_id": current_user.id, "type": "deletion", "status": "pending", "requested_at": now}}, upsert=True,
    )
    await db.users.update_one({"id": current_user.id}, {"$set": {"is_active": False, "deletion_requested_at": now}})
    await db.user_sessions.delete_many({"user_id": current_user.id})
    await db.telegram_bindings.update_many({"user_id": current_user.id}, {"$set": {"active": False, "unlinked_at": now}})
    response = JSONResponse({"accepted": True, "status": "pending_review"}, status_code=202)
    response.delete_cookie("session_token", path="/")
    return response


@api_router.post("/auth/logout")
async def logout(session_token: Optional[str] = Cookie(default=None)):
    if session_token:
        await db.user_sessions.delete_many({"$or": [
            {"session_token_hash": token_digest(session_token)}, {"session_token": session_token}
        ]})
    response = JSONResponse({"success": True})
    response.delete_cookie("session_token", path="/")
    return response


@api_router.post("/users/{user_id}/password-reset-link")
async def create_password_reset_link(user_id: str, current_user: User = Depends(get_current_user)):
    if current_user.user_type != UserType.ADMIN:
        raise HTTPException(status_code=403, detail="Solo administracion puede generar enlaces de recuperacion")
    target = await db.users.find_one({"id": user_id, "is_active": True}, {"_id": 0, "password_hash": 0})
    if not target:
        raise HTTPException(status_code=404, detail="Usuario no encontrado")
    raw_token = secrets.token_urlsafe(48)
    now = datetime.now(timezone.utc)
    await db.password_reset_tokens.update_many(
        {"user_id": user_id, "used_at": None}, {"$set": {"used_at": now, "revoked": True}},
    )
    await db.password_reset_tokens.insert_one({
        "id": str(uuid.uuid4()), "user_id": user_id, "token_hash": token_digest(raw_token),
        "created_by": current_user.id, "created_at": now,
        "expires_at": now + timedelta(minutes=30), "used_at": None,
    })
    await log_activity(current_user, "Generacion de enlace de recuperacion", target, {"expires_in_minutes": 30})
    return {"reset_url": f"{PUBLIC_APP_URL.rstrip('/')}/reset-password?token={raw_token}", "expires_in_minutes": 30}


@api_router.post("/auth/reset-password")
async def reset_password(request: ResetPasswordRequest):
    now = datetime.now(timezone.utc)
    token = await db.password_reset_tokens.find_one_and_update(
        {"token_hash": token_digest(request.token), "used_at": None, "expires_at": {"$gt": now}},
        {"$set": {"used_at": now}}, return_document=ReturnDocument.AFTER,
    )
    if not token:
        raise HTTPException(status_code=400, detail="El enlace no es valido o ha caducado")
    await db.users.update_one({"id": token["user_id"]}, {"$set": {"password_hash": password_hash(request.password)}})
    await db.user_sessions.delete_many({"user_id": token["user_id"]})
    return {"success": True}


@api_router.get("/users", response_model=List[User])
async def list_users(current_user: User = Depends(get_current_user)):
    if current_user.user_type not in {UserType.ADMIN, UserType.JUNTA_DIRECTIVA}:
        raise HTTPException(status_code=403, detail="Sin permisos para consultar usuarios")
    return [User(**item) for item in await db.users.find({}, {"_id": 0, "password_hash": 0}).to_list(5000)]


@api_router.post("/users", response_model=User)
async def create_user(request: UserCreate, current_user: User = Depends(get_current_user)):
    if current_user.user_type != UserType.ADMIN:
        raise HTTPException(status_code=403, detail="Sin permisos para crear usuarios")
    if request.user_type == UserType.JUNTA_DIRECTIVA:
        raise HTTPException(status_code=400, detail="Añade las personas de Junta desde su configuración")
    if await db.users.find_one({"email": request.email.lower()}):
        raise HTTPException(status_code=409, detail="El correo ya está registrado")
    if request.user_type in {UserType.REPRESENTANTE, UserType.UNIVERSIDAD}:
        if not request.university_id or not await db.universities.find_one({"id": request.university_id, "is_active": True}):
            raise HTTPException(status_code=400, detail="Este perfil requiere una universidad activa")
    user = User(email=request.email.lower(), name=request.name.strip(), user_type=request.user_type, university_id=request.university_id)
    payload = serialize(user)
    payload["password_hash"] = password_hash(request.password)
    await db.users.insert_one(payload)
    await log_activity(current_user, "Alta de usuario", serialize(user))
    return user


@api_router.put("/users/{user_id}", response_model=User)
async def update_user(user_id: str, request: UpdateUserRequest, current_user: User = Depends(get_current_user)):
    if current_user.user_type != UserType.ADMIN:
        raise HTTPException(status_code=403, detail="Sin permisos para editar usuarios")
    stored = await db.users.find_one({"id": user_id})
    if not stored:
        raise HTTPException(status_code=404, detail="Usuario no encontrado")
    if request.user_type == UserType.JUNTA_DIRECTIVA:
        raise HTTPException(status_code=400, detail="Gestiona este rol desde la Junta Directiva")
    if stored.get("user_type") == UserType.JUNTA_DIRECTIVA.value and request.user_type and request.user_type != UserType.JUNTA_DIRECTIVA:
        raise HTTPException(status_code=400, detail="Actualiza la composición completa de Junta Directiva")
    if stored.get("user_type") == UserType.JUNTA_DIRECTIVA.value and request.is_active is False:
        raise HTTPException(status_code=400, detail="No se puede desactivar una persona de Junta; actualiza primero su composición")
    changes = request.model_dump(exclude_none=True, mode="json")
    if changes:
        await db.users.update_one({"id": user_id}, {"$set": changes})
        if changes.get("is_active") is False:
            await db.user_sessions.delete_many({"user_id": user_id})
        await log_activity(current_user, "Actualización de usuario", stored, {
            "changed_fields": sorted(changes)
        })
    return User(**await db.users.find_one({"id": user_id}, {"_id": 0, "password_hash": 0}))


@api_router.delete("/users/{user_id}", status_code=204)
async def delete_user(user_id: str, current_user: User = Depends(get_current_user)):
    if current_user.user_type != UserType.ADMIN or current_user.id == user_id:
        raise HTTPException(status_code=403, detail="No puedes eliminar este usuario")
    stored = await db.users.find_one({"id": user_id})
    if not stored:
        raise HTTPException(status_code=404, detail="Usuario no encontrado")
    if stored.get("user_type") == UserType.JUNTA_DIRECTIVA.value:
        raise HTTPException(status_code=400, detail="Actualiza primero la composición de Junta Directiva")
    await db.user_sessions.delete_many({"user_id": user_id})
    await db.user_progress.delete_many({"user_id": user_id})
    await db.enrollments.delete_many({"user_id": user_id})
    await db.content_assignments.update_many({}, {"$pull": {"assigned_to_user_ids": user_id}})
    await db.synchronous_sessions.update_many({}, {"$pull": {"participant_user_ids": user_id}})
    await db.users.delete_one({"id": user_id})


@api_router.get("/board", response_model=List[User])
async def get_board(current_user: User = Depends(get_current_user)):
    return [User(**item) for item in await db.users.find({"user_type": UserType.JUNTA_DIRECTIVA.value}, {"_id": 0, "password_hash": 0}).to_list(7)]


@api_router.put("/board", response_model=List[User])
async def update_board(request: BoardUpdateRequest, current_user: User = Depends(get_current_user)):
    if current_user.user_type != UserType.ADMIN:
        raise HTTPException(status_code=403, detail="Solo administración puede configurar la Junta Directiva")
    try:
        validate_board_members(request.members)
    except ValueError as error:
        raise HTTPException(status_code=400, detail=str(error)) from error
    ids = [item.user_id for item in request.members]
    if await db.users.count_documents({"id": {"$in": ids}, "is_active": True, "user_type": {"$ne": UserType.ADMIN.value}}) != len(ids):
        raise HTTPException(status_code=400, detail="La Junta solo puede incluir personas activas que no sean administradoras")
    previous = await db.users.find({"user_type": UserType.JUNTA_DIRECTIVA.value}, {"id": 1}).to_list(7)
    removed_ids = [item["id"] for item in previous if item["id"] not in ids]
    if removed_ids and await db.vocalias.count_documents({"board_member_id": {"$in": removed_ids}}):
        raise HTTPException(status_code=400, detail="Reasigna las Vocalías de las personas que salen de Junta")
    if removed_ids:
        await db.users.update_many({"id": {"$in": removed_ids}}, {"$set": {"user_type": UserType.REPRESENTANTE.value, "board_position": None}})
    for member in request.members:
        await db.users.update_one({"id": member.user_id}, {"$set": {"user_type": UserType.JUNTA_DIRECTIVA.value, "board_position": member.position.value}})
    await log_activity(current_user, "Actualización de Junta Directiva", {"id": "board", "name": "Junta Directiva"})
    return await get_board(current_user)


async def validate_board_owner(user_id: str):
    if not await db.users.find_one({"id": user_id, "user_type": UserType.JUNTA_DIRECTIVA.value, "is_active": True}):
        raise HTTPException(status_code=400, detail="La Vocalía debe asociarse a una persona activa de Junta Directiva")


@api_router.get("/vocalias", response_model=List[Vocalia])
async def list_vocalias(current_user: User = Depends(get_current_user)):
    return [Vocalia(**item) for item in await db.vocalias.find({}, {"_id": 0}).to_list(1000)]


@api_router.post("/vocalias", response_model=Vocalia)
async def create_vocalia(request: VocaliaCreate, current_user: User = Depends(get_current_user)):
    if current_user.user_type != UserType.ADMIN:
        raise HTTPException(status_code=403, detail="Solo administración puede crear Vocalías")
    await validate_board_owner(request.board_member_id)
    if await db.vocalias.find_one({"name": {"$regex": f"^{re.escape(request.name)}$", "$options": "i"}}):
        raise HTTPException(status_code=409, detail="Ya existe una Vocalía con ese nombre")
    vocalia = Vocalia(**request.model_dump())
    await db.vocalias.insert_one(serialize(vocalia))
    return vocalia


@api_router.put("/vocalias/{vocalia_id}", response_model=Vocalia)
async def update_vocalia(vocalia_id: str, request: VocaliaCreate, current_user: User = Depends(get_current_user)):
    if current_user.user_type != UserType.ADMIN:
        raise HTTPException(status_code=403, detail="Solo administración puede editar Vocalías")
    await validate_board_owner(request.board_member_id)
    if not (await db.vocalias.update_one({"id": vocalia_id}, {"$set": request.model_dump(mode="json")})).matched_count:
        raise HTTPException(status_code=404, detail="Vocalía no encontrada")
    return Vocalia(**await db.vocalias.find_one({"id": vocalia_id}, {"_id": 0}))


@api_router.delete("/vocalias/{vocalia_id}")
async def delete_vocalia(vocalia_id: str, current_user: User = Depends(get_current_user)):
    if current_user.user_type != UserType.ADMIN:
        raise HTTPException(status_code=403, detail="Solo administración puede eliminar Vocalías")
    await db.users.update_many({}, {"$pull": {"vocalia_ids": vocalia_id}})
    if not (await db.vocalias.delete_one({"id": vocalia_id})).deleted_count:
        raise HTTPException(status_code=404, detail="Vocalía no encontrada")
    return {"message": "Vocalía eliminada"}


@api_router.put("/vocalias/{vocalia_id}/members")
async def update_vocalia_members(vocalia_id: str, request: AssignUsersRequest, current_user: User = Depends(get_current_user)):
    vocalia = await db.vocalias.find_one({"id": vocalia_id})
    if not vocalia:
        raise HTTPException(status_code=404, detail="Vocalía no encontrada")
    can_manage = current_user.user_type == UserType.ADMIN or (current_user.user_type == UserType.JUNTA_DIRECTIVA and vocalia["board_member_id"] == current_user.id)
    if not can_manage:
        raise HTTPException(status_code=403, detail="No puedes gestionar esta Vocalía")
    if request.user_ids and await db.users.count_documents({"id": {"$in": request.user_ids}, "is_active": True}) != len(set(request.user_ids)):
        raise HTTPException(status_code=400, detail="La selección contiene usuarios inexistentes o inactivos")
    await db.users.update_many({"vocalia_ids": vocalia_id}, {"$pull": {"vocalia_ids": vocalia_id}})
    if request.user_ids:
        await db.users.update_many({"id": {"$in": request.user_ids}}, {"$addToSet": {"vocalia_ids": vocalia_id}})
    return {"message": "Miembros actualizados", "count": len(request.user_ids)}


@api_router.get("/universities", response_model=List[University])
async def get_universities():
    return [University(**item) for item in await db.universities.find({}, {"_id": 0}).sort("name", 1).to_list(1000)]


@api_router.post("/universities/import-ritsi")
async def sync_ritsi_universities(current_user: User = Depends(get_current_user)):
    if current_user.user_type != UserType.ADMIN:
        raise HTTPException(status_code=403, detail="Solo administración puede sincronizar universidades")
    try:
        result = await import_ritsi_universities()
    except requests.RequestException as error:
        raise HTTPException(status_code=502, detail=f"No se pudo consultar el listado oficial: {error}") from error
    await log_activity(current_user, "Sincronización de universidades", {"id": "ritsi-universities", "name": "Socios de RITSI"}, result)
    return result


@api_router.post("/universities", response_model=University)
async def create_university(request: UniversityCreate, current_user: User = Depends(get_current_user)):
    if current_user.user_type != UserType.ADMIN:
        raise HTTPException(status_code=403, detail="Sin permisos")
    university = University(**request.model_dump(), is_active=True, is_ritsi_member=True)
    await db.universities.insert_one(serialize(university))
    return university


@api_router.put("/universities/{university_id}", response_model=University)
async def update_university(university_id: str, request: UniversityCreate, current_user: User = Depends(get_current_user)):
    if current_user.user_type != UserType.ADMIN:
        raise HTTPException(status_code=403, detail="Sin permisos")
    if not (await db.universities.update_one({"id": university_id}, {"$set": request.model_dump()})).matched_count:
        raise HTTPException(status_code=404, detail="Universidad no encontrada")
    return University(**await db.universities.find_one({"id": university_id}, {"_id": 0}))


@api_router.put("/universities/{university_id}/status", response_model=University)
async def update_university_status(university_id: str, request: UpdateStatusRequest, current_user: User = Depends(get_current_user)):
    if current_user.user_type != UserType.ADMIN:
        raise HTTPException(status_code=403, detail="Solo administración puede cambiar el estado de una universidad")
    if not (await db.universities.update_one({"id": university_id}, {"$set": {
        "is_active": request.is_active, "is_ritsi_member": request.is_active
    }})).matched_count:
        raise HTTPException(status_code=404, detail="Universidad no encontrada")
    return University(**await db.universities.find_one({"id": university_id}, {"_id": 0}))


@api_router.delete("/universities/{university_id}", status_code=204)
async def delete_university(university_id: str, current_user: User = Depends(get_current_user)):
    if current_user.user_type != UserType.ADMIN:
        raise HTTPException(status_code=403, detail="Sin permisos")
    if await db.users.count_documents({"university_id": university_id}):
        raise HTTPException(status_code=400, detail="La universidad tiene usuarios asociados")
    await db.universities.delete_one({"id": university_id})


@api_router.get("/categories", response_model=List[Category])
async def get_categories(current_user: User = Depends(get_current_user)):
    return [Category(**item) for item in await db.categories.find({}, {"_id": 0}).to_list(1000)]


@api_router.post("/categories", response_model=Category)
async def create_category(request: CategoryCreate, current_user: User = Depends(get_current_user)):
    if current_user.user_type not in {UserType.ADMIN, UserType.JUNTA_DIRECTIVA}:
        raise HTTPException(status_code=403, detail="Sin permisos")
    category = Category(name=request.name.strip())
    await db.categories.insert_one(serialize(category))
    return category


@api_router.delete("/categories/{category_id}")
async def delete_category(category_id: str, current_user: User = Depends(get_current_user)):
    if current_user.user_type != UserType.ADMIN:
        raise HTTPException(status_code=403, detail="Sin permisos")
    await db.training_contents.update_many({}, {"$pull": {"category_ids": category_id}})
    await db.categories.delete_one({"id": category_id})
    return {"message": "Categoría eliminada"}


def assignment_access_query(user: User) -> Dict[str, Any]:
    clauses: List[Dict[str, Any]] = [{"assigned_to_user_ids": user.id}]
    if user.user_type == UserType.REPRESENTANTE:
        clauses.append({"assigned_to_all_representatives": True})
    return {"$or": clauses}


async def user_has_content_access(user: User, content: Dict[str, Any]) -> bool:
    if user.user_type in {UserType.ADMIN, UserType.JUNTA_DIRECTIVA}:
        return True
    is_published = content.get("status") == ContentStatus.PUBLISHED.value
    if user.user_type == UserType.FORMADOR and (content.get("created_by") == user.id or is_published):
        return True
    if not is_published:
        return False
    if content.get("is_public", False):
        return True
    if await db.enrollments.find_one({
        "content_id": content["id"], "user_id": user.id,
        "status": {"$ne": EnrollmentStatus.CANCELLED.value},
    }):
        return True
    return bool(await db.content_assignments.find_one({
        "content_id": content["id"], **assignment_access_query(user)
    }))


async def require_content_access(user: User, content_id: str) -> Dict[str, Any]:
    content = await db.training_contents.find_one({"id": content_id}, {"_id": 0})
    if not content:
        raise HTTPException(status_code=404, detail="Formación no encontrada")
    if not await user_has_content_access(user, content):
        raise HTTPException(status_code=403, detail="No tienes acceso a esta formación")
    return content


async def upsert_enrollment(
    user_id: str,
    content_id: str,
    assigned_by: str,
    source: EnrollmentSource,
    assignment_id: Optional[str] = None,
) -> Dict[str, Any]:
    now = datetime.now(timezone.utc)
    enrollment = Enrollment(
        content_id=content_id,
        user_id=user_id,
        assigned_by=assigned_by,
        assignment_ids=[],
        source=source,
    )
    insert_fields = mongo_serialize(enrollment)
    insert_fields["source"] = source.value
    insert_fields["status"] = EnrollmentStatus.ASSIGNED.value
    insert_fields.pop("updated_at", None)
    update: Dict[str, Any] = {
        "$setOnInsert": insert_fields,
        "$set": {"updated_at": now},
    }
    if assignment_id:
        insert_fields.pop("assignment_ids", None)
        update["$addToSet"] = {"assignment_ids": assignment_id}
    updated = await db.enrollments.find_one_and_update(
        {"content_id": content_id, "user_id": user_id},
        update,
        upsert=True,
        return_document=ReturnDocument.AFTER,
    )
    if updated.get("status") == EnrollmentStatus.CANCELLED.value:
        updated = await db.enrollments.find_one_and_update(
            {"_id": updated["_id"]},
            {"$set": {
                "status": EnrollmentStatus.ASSIGNED.value,
                "completed_at": None,
                "assigned_by": assigned_by,
                "source": source.value,
                "updated_at": now,
            }},
            return_document=ReturnDocument.AFTER,
        )
    updated.pop("_id", None)
    return updated


async def ensure_legacy_enrollments_for_user(user: User) -> None:
    assignments = await db.content_assignments.find(assignment_access_query(user), {"_id": 0}).to_list(5000)
    for assignment in assignments:
        source = (EnrollmentSource.GLOBAL_ASSIGNMENT if assignment.get("assigned_to_all_representatives")
                  else EnrollmentSource.ASSIGNMENT)
        await upsert_enrollment(user.id, assignment["content_id"], assignment["assigned_by"], source, assignment["id"])


async def update_enrollment_from_progress(user: User, content_id: str, completed: bool) -> None:
    # Preserve assignment provenance so revoking that assignment also revokes
    # access. SELF_ENROLLED is reserved for public catalogue content.
    await ensure_legacy_enrollments_for_user(user)
    enrollment = await upsert_enrollment(
        user.id, content_id, user.id, EnrollmentSource.SELF_ENROLLED,
    )
    previous_status = enrollment.get("status", EnrollmentStatus.ASSIGNED.value)
    now = datetime.now(timezone.utc)
    next_status = EnrollmentStatus.COMPLETED.value if completed else EnrollmentStatus.IN_PROGRESS.value
    changes: Dict[str, Any] = {
        "status": next_status,
        "updated_at": now,
    }
    if not enrollment.get("started_at"):
        changes["started_at"] = now
    changes["completed_at"] = now if completed else None
    await db.enrollments.update_one({"id": enrollment["id"]}, {"$set": changes})
    content = None
    if completed:
        content = await db.training_contents.find_one({"id": content_id}, {"_id": 0})
        if content:
            await db.certificates.update_one(
                {"user_id": user.id, "content_id": content_id},
                {"$setOnInsert": {
                    "id": str(uuid.uuid4()), "verification_code": secrets.token_urlsafe(16),
                    "user_id": user.id, "user_name": user.name, "content_id": content_id,
                    "content_title": content["title"], "issued_at": now, "status": "valid",
                }}, upsert=True,
            )
    if previous_status != next_status:
        if content is None:
            content = await db.training_contents.find_one({"id": content_id}, {"_id": 0})
        event_type = "TrainingCompleted" if completed else "TrainingStarted"
        await enqueue_event(event_type, "Enrollment", enrollment["id"], {
            "content_id": content_id,
            "content_title": content.get("title", "Formación") if content else "Formación",
            "user_ids": [user.id],
        }, f"enrollment:{enrollment['id']}:{next_status}")


@api_router.post("/content/import-sheet")
async def import_sheet(request: ImportFormationsRequest, current_user: User = Depends(get_current_user)):
    if current_user.user_type != UserType.ADMIN:
        raise HTTPException(status_code=403, detail="Solo administración puede importar el catálogo")
    try:
        result = await import_formations(request.spreadsheet_id, current_user.id, request.publish, request.dry_run)
    except requests.RequestException as error:
        raise HTTPException(status_code=502, detail=f"No se pudo descargar la hoja: {error}") from error
    await log_activity(current_user, "Importación del catálogo", {"id": request.spreadsheet_id, "name": "Google Sheets"}, result)
    return result


@api_router.get("/content", response_model=List[TrainingContentView])
async def get_content(current_user: User = Depends(get_current_user)):
    if current_user.user_type in {UserType.ADMIN, UserType.JUNTA_DIRECTIVA}:
        query: Dict[str, Any] = {}
    elif current_user.user_type == UserType.FORMADOR:
        query = {"$or": [{"created_by": current_user.id}, {"status": ContentStatus.PUBLISHED.value}]}
    else:
        assignments = await db.content_assignments.find(assignment_access_query(current_user), {"content_id": 1}).to_list(5000)
        enrollments = await db.enrollments.find({
            "user_id": current_user.id, "status": {"$ne": EnrollmentStatus.CANCELLED.value}
        }, {"content_id": 1}).to_list(5000)
        assigned_content_ids = list({
            *(item["content_id"] for item in assignments),
            *(item["content_id"] for item in enrollments),
        })
        query = {"status": ContentStatus.PUBLISHED.value, "$or": [
            {"is_public": True}, {"id": {"$in": assigned_content_ids}}
        ]}
    items = await db.training_contents.find(query, {"_id": 0}).sort([("training_date", -1), ("title", 1)]).to_list(10000)
    return [TrainingContentView(**item) for item in items]


@api_router.post("/content", response_model=TrainingContent)
async def create_content(request: TrainingContentCreate, current_user: User = Depends(get_current_user)):
    if current_user.user_type not in {UserType.ADMIN, UserType.JUNTA_DIRECTIVA, UserType.FORMADOR}:
        raise HTTPException(status_code=403, detail="Sin permisos para crear formaciones")
    content = TrainingContent(**request.model_dump(), created_by=current_user.id,
        status=ContentStatus.PUBLISHED if current_user.user_type in {UserType.ADMIN, UserType.JUNTA_DIRECTIVA} else ContentStatus.PENDING)
    await db.training_contents.insert_one(serialize(content))
    return content


@api_router.get("/content/{content_id}", response_model=TrainingContentView)
async def get_content_item(content_id: str, current_user: User = Depends(get_current_user)):
    return TrainingContentView(**await require_content_access(current_user, content_id))


@api_router.get("/content/{content_id}/editor", response_model=TrainingContent)
async def get_content_editor(content_id: str, current_user: User = Depends(get_current_user)):
    content = await db.training_contents.find_one({"id": content_id}, {"_id": 0})
    if not content:
        raise HTTPException(status_code=404, detail="Formacion no encontrada")
    reviewer = current_user.user_type in {UserType.ADMIN, UserType.JUNTA_DIRECTIVA}
    if not reviewer and not (current_user.user_type == UserType.FORMADOR and content.get("created_by") == current_user.id):
        raise HTTPException(status_code=403, detail="No puedes editar esta formacion")
    return TrainingContent(**content)


@api_router.post("/content/{content_id}/approve", response_model=TrainingContent)
async def approve_content(content_id: str, current_user: User = Depends(get_current_user)):
    if current_user.user_type not in {UserType.ADMIN, UserType.JUNTA_DIRECTIVA}:
        raise HTTPException(status_code=403, detail="Sin permisos")
    if not (await db.training_contents.update_one({"id": content_id}, {"$set": {"status": ContentStatus.PUBLISHED.value}})).matched_count:
        raise HTTPException(status_code=404, detail="Formación no encontrada")
    return TrainingContent(**await db.training_contents.find_one({"id": content_id}, {"_id": 0}))


@api_router.patch("/content/{content_id}", response_model=TrainingContent)
async def update_content(content_id: str, request: ContentUpdateRequest, current_user: User = Depends(get_current_user)):
    content = await db.training_contents.find_one({"id": content_id})
    if not content:
        raise HTTPException(status_code=404, detail="Formacion no encontrada")
    reviewer = current_user.user_type in {UserType.ADMIN, UserType.JUNTA_DIRECTIVA}
    if not reviewer and not (current_user.user_type == UserType.FORMADOR and content.get("created_by") == current_user.id):
        raise HTTPException(status_code=403, detail="No puedes editar esta formacion")
    if content.get("status") == ContentStatus.PUBLISHED.value and not reviewer:
        raise HTTPException(status_code=409, detail="Solicita cambios antes de editar una formacion publicada")
    changes = request.model_dump(exclude_unset=True, mode="json")
    if changes:
        changes["updated_at"] = datetime.now(timezone.utc)
        await db.training_contents.update_one({"_id": content["_id"]}, {"$set": changes})
    return TrainingContent(**await db.training_contents.find_one({"id": content_id}, {"_id": 0}))


@api_router.post("/content/{content_id}/submit-review", response_model=TrainingContent)
async def submit_content_review(content_id: str, current_user: User = Depends(get_current_user)):
    result = await db.training_contents.find_one_and_update(
        {"id": content_id, "created_by": current_user.id, "status": {"$in": [ContentStatus.DRAFT.value, ContentStatus.CHANGES_REQUESTED.value]}},
        {"$set": {"status": ContentStatus.PENDING.value, "review_note": None, "updated_at": datetime.now(timezone.utc)}},
        return_document=ReturnDocument.AFTER,
    )
    if not result:
        raise HTTPException(status_code=409, detail="La formacion no puede enviarse a revision")
    return TrainingContent(**result)


@api_router.post("/content/{content_id}/review", response_model=TrainingContent)
async def review_content(content_id: str, request: ReviewRequest, current_user: User = Depends(get_current_user)):
    if current_user.user_type not in {UserType.ADMIN, UserType.JUNTA_DIRECTIVA}:
        raise HTTPException(status_code=403, detail="Sin permisos de revision")
    statuses = {
        "publish": ContentStatus.PUBLISHED.value,
        "request_changes": ContentStatus.CHANGES_REQUESTED.value,
        "archive": ContentStatus.ARCHIVED.value,
    }
    result = await db.training_contents.find_one_and_update(
        {"id": content_id}, {"$set": {"status": statuses[request.action], "review_note": request.note, "reviewed_by": current_user.id, "reviewed_at": datetime.now(timezone.utc)}},
        return_document=ReturnDocument.AFTER,
    )
    if not result:
        raise HTTPException(status_code=404, detail="Formacion no encontrada")
    return TrainingContent(**result)


@api_router.delete("/content/{content_id}")
async def delete_content(content_id: str, current_user: User = Depends(get_current_user)):
    if current_user.user_type not in {UserType.ADMIN, UserType.JUNTA_DIRECTIVA}:
        raise HTTPException(status_code=403, detail="Sin permisos")
    affected_sessions = await db.synchronous_sessions.find({"content_id": content_id, "status": {"$ne": SynchronousSessionStatus.CANCELLED.value}}, {"id": 1, "starts_at": 1, "participant_user_ids": 1}).to_list(5000)
    inherited_attendees = [item["user_id"] for item in await db.enrollments.find(
        {"content_id": content_id, "status": {"$ne": EnrollmentStatus.CANCELLED.value}}, {"user_id": 1}
    ).to_list(10000)]
    await db.training_contents.delete_one({"id": content_id})
    await db.content_assignments.delete_many({"content_id": content_id})
    await db.enrollments.delete_many({"content_id": content_id})
    await db.user_progress.delete_many({"content_id": content_id})
    await db.synchronous_sessions.update_many(
        {"content_id": content_id},
        {"$set": {"status": SynchronousSessionStatus.CANCELLED.value, "updated_at": datetime.now(timezone.utc)}},
    )
    for session in affected_sessions:
        await enqueue_event("SessionCancelled", "SynchronousSession", session["id"], {
            "session_id": session["id"], "starts_at": session["starts_at"],
            "user_ids": session.get("participant_user_ids") or inherited_attendees,
        }, f"session:{session['id']}:content-deleted")
    return {"message": "Formación eliminada"}


@api_router.get("/assignments/eligible-users", response_model=List[User])
async def assignment_eligible_users(current_user: User = Depends(get_current_user)):
    if current_user.user_type in {UserType.ADMIN, UserType.JUNTA_DIRECTIVA}:
        query: Dict[str, Any] = {"is_active": True, "user_type": {"$ne": UserType.ADMIN.value}}
    elif current_user.user_type == UserType.UNIVERSIDAD and current_user.university_id:
        query = {
            "is_active": True,
            "university_id": current_user.university_id,
            "user_type": {"$ne": UserType.ADMIN.value},
        }
    else:
        raise HTTPException(status_code=403, detail="Sin permisos para consultar destinatarios")
    items = await db.users.find(query, {"_id": 0, "password_hash": 0}).sort("name", 1).to_list(5000)
    return [User(**item) for item in items]


@api_router.get("/assignments", response_model=List[ContentAssignment])
async def list_assignments(current_user: User = Depends(get_current_user)):
    if current_user.user_type in {UserType.ADMIN, UserType.JUNTA_DIRECTIVA}:
        query: Dict[str, Any] = {}
    elif current_user.user_type == UserType.UNIVERSIDAD:
        query = {"assigned_by": current_user.id}
    else:
        raise HTTPException(status_code=403, detail="Sin permisos para consultar asignaciones")
    return [ContentAssignment(**item) for item in await db.content_assignments.find(
        query, {"_id": 0}
    ).sort("created_at", -1).to_list(5000)]


@api_router.get("/assignments/me", response_model=List[Enrollment])
async def my_assignments(current_user: User = Depends(get_current_user)):
    await ensure_legacy_enrollments_for_user(current_user)
    items = await db.enrollments.find({
        "user_id": current_user.id,
        "status": {"$ne": EnrollmentStatus.CANCELLED.value},
    }, {"_id": 0}).sort("created_at", -1).to_list(5000)
    return [Enrollment(**item) for item in items]


@api_router.post("/assignments", response_model=ContentAssignment, status_code=201)
async def assign_content(request: AssignContentRequest, current_user: User = Depends(get_current_user)):
    if current_user.user_type not in {UserType.ADMIN, UserType.JUNTA_DIRECTIVA, UserType.UNIVERSIDAD}:
        raise HTTPException(status_code=403, detail="Sin permisos")
    content = await db.training_contents.find_one({"id": request.content_id}, {"_id": 0})
    if not content:
        raise HTTPException(status_code=404, detail="Formación no encontrada")
    if content.get("status") != ContentStatus.PUBLISHED.value:
        raise HTTPException(status_code=409, detail="Solo se puede asignar una formación publicada")

    user_ids = list(dict.fromkeys(request.user_ids))
    if not user_ids and not request.assign_to_all_representatives:
        raise HTTPException(status_code=422, detail="Selecciona al menos una persona o una audiencia global")
    if current_user.user_type == UserType.UNIVERSIDAD and request.assign_to_all_representatives:
        raise HTTPException(status_code=403, detail="Una universidad no puede realizar asignaciones globales")

    recipient_query: Dict[str, Any] = {
        "id": {"$in": user_ids}, "is_active": True, "user_type": {"$ne": UserType.ADMIN.value}
    }
    if current_user.user_type == UserType.UNIVERSIDAD:
        if not current_user.university_id:
            raise HTTPException(status_code=403, detail="La cuenta no tiene una universidad asociada")
        recipient_query["university_id"] = current_user.university_id
    if user_ids and await db.users.count_documents(recipient_query) != len(user_ids):
        if current_user.user_type == UserType.UNIVERSIDAD:
            raise HTTPException(status_code=403, detail="Solo puedes asignar personas activas de tu universidad")
        raise HTTPException(status_code=400, detail="La selección contiene usuarios inexistentes, inactivos o no asignables")

    assignment = ContentAssignment(
        content_id=request.content_id,
        assigned_to_user_ids=user_ids,
        assigned_to_all_representatives=request.assign_to_all_representatives,
        assigned_by=current_user.id,
    )
    await db.content_assignments.insert_one(mongo_serialize(assignment))
    for user_id in user_ids:
        await upsert_enrollment(
            user_id, request.content_id, current_user.id, EnrollmentSource.ASSIGNMENT, assignment.id,
        )
    await enqueue_event("TrainingAssigned", "ContentAssignment", assignment.id, {
        "content_id": request.content_id, "content_title": content["title"],
        "user_ids": user_ids, "all_representatives": request.assign_to_all_representatives,
    }, f"assignment:{assignment.id}:created")
    await log_activity(current_user, "Asignación de formación", content, {
        "assignment_id": assignment.id,
        "recipient_count": len(user_ids),
        "all_representatives": assignment.assigned_to_all_representatives,
    })
    return assignment


@api_router.delete("/assignments/{assignment_id}", status_code=204)
async def unassign_content(assignment_id: str, current_user: User = Depends(get_current_user)):
    assignment = await db.content_assignments.find_one({"id": assignment_id})
    if not assignment:
        raise HTTPException(status_code=404, detail="Asignación no encontrada")
    can_revoke = current_user.user_type in {UserType.ADMIN, UserType.JUNTA_DIRECTIVA}
    if current_user.user_type == UserType.UNIVERSIDAD:
        can_revoke = assignment.get("assigned_by") == current_user.id
    if not can_revoke:
        raise HTTPException(status_code=403, detail="No puedes revocar esta asignación")

    affected_enrollments = await db.enrollments.find(
        {"assignment_ids": assignment_id}, {"_id": 1}
    ).to_list(5000)
    affected_ids = [item["_id"] for item in affected_enrollments]
    await db.content_assignments.delete_one({"_id": assignment["_id"]})
    await db.enrollments.update_many(
        {"_id": {"$in": affected_ids}},
        {"$pull": {"assignment_ids": assignment_id}, "$set": {"updated_at": datetime.now(timezone.utc)}},
    )
    await db.enrollments.update_many({
        "_id": {"$in": affected_ids},
        "assignment_ids": {"$size": 0},
        "source": {"$in": [EnrollmentSource.ASSIGNMENT.value, EnrollmentSource.GLOBAL_ASSIGNMENT.value]},
    }, {"$set": {"status": EnrollmentStatus.CANCELLED.value, "updated_at": datetime.now(timezone.utc)}})
    await log_activity(current_user, "Revocación de asignación", {
        "id": assignment_id, "name": "Asignación de formación"
    }, {"content_id": assignment["content_id"]})


@api_router.get("/learning-paths")
async def list_learning_paths(current_user: User = Depends(get_current_user)):
    query = {} if current_user.user_type in {UserType.ADMIN, UserType.JUNTA_DIRECTIVA} else {"is_active": True}
    return await db.learning_paths.find(query, {"_id": 0}).sort("title", 1).to_list(1000)


@api_router.post("/learning-paths", status_code=201)
async def create_learning_path(request: LearningPathRequest, current_user: User = Depends(get_current_user)):
    if current_user.user_type not in {UserType.ADMIN, UserType.JUNTA_DIRECTIVA}:
        raise HTTPException(status_code=403, detail="Sin permisos para crear itinerarios")
    content_ids = list(dict.fromkeys(request.content_ids))
    if await db.training_contents.count_documents({"id": {"$in": content_ids}, "status": ContentStatus.PUBLISHED.value}) != len(content_ids):
        raise HTTPException(status_code=422, detail="El itinerario solo puede contener formaciones publicadas")
    now = datetime.now(timezone.utc)
    path = {"id": str(uuid.uuid4()), **request.model_dump(), "content_ids": content_ids, "created_by": current_user.id, "created_at": now, "updated_at": now}
    await db.learning_paths.insert_one(path)
    path.pop("_id", None)
    return path


@api_router.put("/learning-paths/{path_id}")
async def update_learning_path(path_id: str, request: LearningPathRequest, current_user: User = Depends(get_current_user)):
    if current_user.user_type not in {UserType.ADMIN, UserType.JUNTA_DIRECTIVA}:
        raise HTTPException(status_code=403, detail="Sin permisos para editar itinerarios")
    content_ids = list(dict.fromkeys(request.content_ids))
    if await db.training_contents.count_documents({"id": {"$in": content_ids}, "status": ContentStatus.PUBLISHED.value}) != len(content_ids):
        raise HTTPException(status_code=422, detail="El itinerario solo puede contener formaciones publicadas")
    result = await db.learning_paths.find_one_and_update(
        {"id": path_id}, {"$set": {**request.model_dump(), "content_ids": content_ids, "updated_at": datetime.now(timezone.utc)}},
        return_document=ReturnDocument.AFTER,
    )
    if not result:
        raise HTTPException(status_code=404, detail="Itinerario no encontrado")
    result.pop("_id", None)
    return result


@api_router.post("/learning-paths/{path_id}/assign", status_code=201)
async def assign_learning_path(path_id: str, request: LearningPathAssignmentRequest, current_user: User = Depends(get_current_user)):
    if current_user.user_type not in {UserType.ADMIN, UserType.JUNTA_DIRECTIVA, UserType.UNIVERSIDAD}:
        raise HTTPException(status_code=403, detail="Sin permisos para asignar itinerarios")
    path = await db.learning_paths.find_one({"id": path_id, "is_active": True}, {"_id": 0})
    if not path:
        raise HTTPException(status_code=404, detail="Itinerario no encontrado")
    user_ids = list(dict.fromkeys(request.user_ids))
    user_query: Dict[str, Any] = {"id": {"$in": user_ids}, "is_active": True, "user_type": {"$ne": UserType.ADMIN.value}}
    if current_user.user_type == UserType.UNIVERSIDAD:
        user_query["university_id"] = current_user.university_id
    if await db.users.count_documents(user_query) != len(user_ids):
        raise HTTPException(status_code=403, detail="Hay destinatarios fuera de tu ambito")
    assignment_id = str(uuid.uuid4())
    now = datetime.now(timezone.utc)
    await db.learning_path_assignments.insert_one({"id": assignment_id, "path_id": path_id, "user_ids": user_ids, "assigned_by": current_user.id, "created_at": now})
    contents = await db.training_contents.find({"id": {"$in": path["content_ids"]}}, {"_id": 0, "id": 1, "title": 1}).to_list(100)
    by_id = {item["id"]: item for item in contents}
    for content_id in path["content_ids"]:
        for user_id in user_ids:
            await upsert_enrollment(user_id, content_id, current_user.id, EnrollmentSource.ASSIGNMENT, f"path:{assignment_id}")
        await enqueue_event("TrainingAssigned", "LearningPathAssignment", assignment_id, {
            "content_id": content_id, "content_title": by_id.get(content_id, {}).get("title", path["title"]), "user_ids": user_ids,
        }, f"path-assignment:{assignment_id}:{content_id}")
    return {"id": assignment_id, "assigned": len(user_ids), "formations": len(path["content_ids"])}


@api_router.get("/learning-paths/me")
async def my_learning_paths(current_user: User = Depends(get_current_user)):
    assignments = await db.learning_path_assignments.find({"user_ids": current_user.id}, {"_id": 0}).to_list(1000)
    path_ids = [item["path_id"] for item in assignments]
    paths = await db.learning_paths.find({"id": {"$in": path_ids}, "is_active": True}, {"_id": 0}).to_list(1000)
    progress = await db.user_progress.find({"user_id": current_user.id}, {"_id": 0, "content_id": 1, "completed": 1}).to_list(10000)
    completed = {item["content_id"] for item in progress if item.get("completed")}
    return [{**path, "completed_count": len(set(path["content_ids"]) & completed), "total_count": len(path["content_ids"]), "completed": set(path["content_ids"]).issubset(completed)} for path in paths]


@api_router.get("/certificates/me")
async def my_certificates(current_user: User = Depends(get_current_user)):
    return await db.certificates.find({"user_id": current_user.id, "status": "valid"}, {"_id": 0}).sort("issued_at", -1).to_list(1000)


@api_router.get("/certificates/verify/{verification_code}")
async def verify_certificate(verification_code: str):
    certificate = await db.certificates.find_one({"verification_code": verification_code, "status": "valid"}, {"_id": 0, "user_id": 0})
    if not certificate:
        raise HTTPException(status_code=404, detail="Certificado no valido")
    return certificate


@api_router.get("/progress/all")
async def get_all_progress(current_user: User = Depends(get_current_user)):
    if current_user.user_type not in {UserType.ADMIN, UserType.JUNTA_DIRECTIVA}:
        raise HTTPException(status_code=403, detail="Sin permisos para consultar el progreso global")
    return await db.user_progress.find({}, {"_id": 0}).to_list(10000)


@api_router.get("/progress/me")
@api_router.get("/progress")
async def get_progress(current_user: User = Depends(get_current_user)):
    return await db.user_progress.find({"user_id": current_user.id}, {"_id": 0}).to_list(10000)


@api_router.get("/progress/me/{content_id}")
async def get_my_content_progress(content_id: str, current_user: User = Depends(get_current_user)):
    await require_content_access(current_user, content_id)
    progress = await db.user_progress.find_one(
        {"user_id": current_user.id, "content_id": content_id}, {"_id": 0}
    )
    return progress or {
        "user_id": current_user.id, "content_id": content_id,
        "files_completed": [], "quizzes_completed": {}, "completed": False,
    }


@api_router.post("/progress/file-completed")
async def mark_file_completed(request: MarkFileCompletedRequest, current_user: User = Depends(get_current_user)):
    content = await require_content_access(current_user, request.content_id)
    if not any(item.get("id") == request.file_id for item in content.get("files", [])):
        raise HTTPException(status_code=404, detail="El recurso no pertenece a esta formación")
    now = datetime.now(timezone.utc)
    progress = await db.user_progress.find_one_and_update(
        {"user_id": current_user.id, "content_id": request.content_id},
        {
            "$setOnInsert": {
                "id": str(uuid.uuid4()), "user_id": current_user.id,
                "content_id": request.content_id, "quizzes_completed": {}, "completed": False,
            },
            "$addToSet": {"files_completed": request.file_id},
            "$set": {"last_updated": now},
        },
        upsert=True,
        return_document=ReturnDocument.AFTER,
    )
    completed = calculate_progress_completion(content, progress)
    await db.user_progress.update_one({"_id": progress["_id"]}, {"$set": {"completed": completed}})
    await update_enrollment_from_progress(current_user, request.content_id, completed)
    progress.pop("_id", None)
    progress["completed"] = completed
    return {"message": "Progreso actualizado", "progress": progress}


@api_router.post("/progress/submit-quiz")
async def submit_quiz(request: SubmitQuizRequest, current_user: User = Depends(get_current_user)):
    content = await require_content_access(current_user, request.content_id)
    quiz = next((item for item in content.get("quizzes", []) if item["id"] == request.quiz_id), None)
    if not quiz:
        raise HTTPException(status_code=404, detail="Cuestionario no encontrado")
    question_ids = {item["id"] for item in quiz.get("questions", [])}
    if set(request.answers) - question_ids:
        raise HTTPException(status_code=422, detail="La entrega contiene preguntas desconocidas")
    for question in quiz.get("questions", []):
        answers = request.answers.get(question["id"], [])
        if len(set(answers)) != len(answers) or any(index < 0 or index >= len(question.get("options", [])) for index in answers):
            raise HTTPException(status_code=422, detail="La entrega contiene respuestas no válidas")

    correct = sum(
        1 for question in quiz.get("questions", [])
        if sorted(request.answers.get(question["id"], [])) == sorted(question.get("correct_answers", []))
    )
    score = (correct / len(quiz["questions"]) * 100) if quiz.get("questions") else 0
    now = datetime.now(timezone.utc)
    result = {
        "score": score,
        "passed": score >= quiz.get("passing_percentage", 70),
        "completed_at": now,
    }
    previous_attempts = await db.assessment_attempts.count_documents({
        "user_id": current_user.id, "content_id": request.content_id, "quiz_id": request.quiz_id,
    })
    await db.assessment_attempts.insert_one({
        "id": str(uuid.uuid4()), "user_id": current_user.id,
        "content_id": request.content_id, "quiz_id": request.quiz_id,
        "attempt_number": previous_attempts + 1, "answers": request.answers,
        "score": score, "passed": result["passed"], "submitted_at": now,
    })
    progress = await db.user_progress.find_one_and_update(
        {"user_id": current_user.id, "content_id": request.content_id},
        {
            "$setOnInsert": {
                "id": str(uuid.uuid4()), "user_id": current_user.id,
                "content_id": request.content_id, "files_completed": [], "completed": False,
            },
            "$set": {f"quizzes_completed.{request.quiz_id}": result, "last_updated": now},
        },
        upsert=True,
        return_document=ReturnDocument.AFTER,
    )
    completed = calculate_progress_completion(content, progress)
    await db.user_progress.update_one({"_id": progress["_id"]}, {"$set": {"completed": completed}})
    await update_enrollment_from_progress(current_user, request.content_id, completed)
    return {**result, "completed": completed}


async def can_manage_session(user: User, session: Dict[str, Any]) -> bool:
    if user.user_type in {UserType.ADMIN, UserType.JUNTA_DIRECTIVA}:
        return True
    if user.user_type != UserType.FORMADOR:
        return False
    if session.get("created_by") == user.id:
        return True
    return bool(await db.training_contents.find_one({
        "id": session["content_id"], "created_by": user.id
    }, {"id": 1}))


async def can_participate_in_session(user: User, session: Dict[str, Any]) -> bool:
    if await can_manage_session(user, session):
        return True
    participants = session.get("participant_user_ids", [])
    if participants:
        return user.id in participants
    return bool(await db.enrollments.find_one({
        "content_id": session["content_id"],
        "user_id": user.id,
        "status": {"$ne": EnrollmentStatus.CANCELLED.value},
    }))


async def require_session_access(user: User, session_id: str) -> Dict[str, Any]:
    session = await db.synchronous_sessions.find_one({"id": session_id})
    if not session:
        raise HTTPException(status_code=404, detail="Sesión no encontrada")
    await ensure_legacy_enrollments_for_user(user)
    if not await can_participate_in_session(user, session):
        raise HTTPException(status_code=403, detail="No tienes acceso a esta sesión")
    return session


async def build_session_view(session: Dict[str, Any], user: User) -> SynchronousSessionView:
    session_data = dict(session)
    session_data.pop("_id", None)
    starts_at = parse_stored_utc_datetime(session_data["starts_at"])
    ends_at = parse_stored_utc_datetime(session_data["ends_at"])
    session_data["starts_at"] = starts_at
    session_data["ends_at"] = ends_at
    session_data["created_at"] = parse_stored_utc_datetime(session_data.get("created_at", starts_at))
    session_data["updated_at"] = parse_stored_utc_datetime(session_data.get("updated_at", starts_at))
    participant_ids = list(session_data.pop("participant_user_ids", []))
    participant_count = len(participant_ids)
    now = datetime.now(timezone.utc)
    persisted_status = SynchronousSessionStatus(
        session_data.get("status", SynchronousSessionStatus.SCHEDULED.value)
    )
    if persisted_status in {SynchronousSessionStatus.SCHEDULED, SynchronousSessionStatus.LIVE}:
        if now >= ends_at:
            session_data["status"] = SynchronousSessionStatus.ENDED
        elif now >= starts_at:
            session_data["status"] = SynchronousSessionStatus.LIVE
    allowed_by_time, reason = session_join_state(session_data, now)
    allowed_identity = await can_participate_in_session(
        user, {**session_data, "participant_user_ids": participant_ids},
    )
    has_meeting = bool(read_meeting_url(session))
    can_join = allowed_by_time and allowed_identity and has_meeting
    if not allowed_identity:
        reason = "No estás inscrito en esta sesión"
    elif allowed_by_time and not has_meeting:
        reason = "El enlace de acceso aún no está configurado"
    return SynchronousSessionView(
        **session_data,
        participant_count=participant_count,
        can_join=can_join,
        join_unavailable_reason=None if can_join else reason,
        join_available_from=starts_at - timedelta(
            minutes=int(session_data.get("join_window_minutes", DEFAULT_JOIN_WINDOW_MINUTES))
        ),
        can_manage=await can_manage_session(user, session),
    )


async def validate_session_participants(
    user_ids: List[str], content: Dict[str, Any], actor: User,
) -> List[str]:
    unique_ids = list(dict.fromkeys(user_ids))
    if not unique_ids:
        return []
    users = await db.users.find({
        "id": {"$in": unique_ids}, "is_active": True, "user_type": {"$ne": UserType.ADMIN.value}
    }, {"_id": 0, "password_hash": 0}).to_list(5000)
    if len(users) != len(unique_ids):
        raise HTTPException(status_code=400, detail="La lista contiene participantes inexistentes o inactivos")
    if actor.user_type not in {UserType.ADMIN, UserType.JUNTA_DIRECTIVA}:
        for item in users:
            if not await user_has_content_access(User(**item), content):
                raise HTTPException(
                    status_code=403,
                    detail="Solo puedes añadir participantes con acceso a la formación",
                )
    return unique_ids


@api_router.get("/sessions", response_model=List[SynchronousSessionView])
async def list_synchronous_sessions(current_user: User = Depends(get_current_user)):
    if current_user.user_type in {UserType.ADMIN, UserType.JUNTA_DIRECTIVA}:
        query: Dict[str, Any] = {}
    elif current_user.user_type == UserType.FORMADOR:
        content_ids = [item["id"] for item in await db.training_contents.find(
            {"created_by": current_user.id}, {"id": 1}
        ).to_list(5000)]
        query = {"$or": [
            {"created_by": current_user.id}, {"content_id": {"$in": content_ids}}
        ]}
    else:
        await ensure_legacy_enrollments_for_user(current_user)
        content_ids = [item["content_id"] for item in await db.enrollments.find({
            "user_id": current_user.id,
            "status": {"$ne": EnrollmentStatus.CANCELLED.value},
        }, {"content_id": 1}).to_list(5000)]
        query = {"$or": [
            {"participant_user_ids": current_user.id},
            {"$and": [
                {"$or": [
                    {"participant_user_ids": {"$size": 0}},
                    {"participant_user_ids": {"$exists": False}},
                ]},
                {"content_id": {"$in": content_ids}},
            ]},
        ]}
    sessions = await db.synchronous_sessions.find(query).sort("starts_at", 1).to_list(5000)
    return [await build_session_view(item, current_user) for item in sessions]


@api_router.post("/sessions", response_model=SynchronousSessionView, status_code=201)
async def create_synchronous_session(
    request: SessionCreateRequest, current_user: User = Depends(get_current_user),
):
    if current_user.user_type not in {UserType.ADMIN, UserType.JUNTA_DIRECTIVA, UserType.FORMADOR}:
        raise HTTPException(status_code=403, detail="Sin permisos para crear sesiones")
    content = await db.training_contents.find_one({"id": request.content_id}, {"_id": 0})
    if not content:
        raise HTTPException(status_code=404, detail="Formación no encontrada")
    if current_user.user_type == UserType.FORMADOR and content.get("created_by") != current_user.id:
        raise HTTPException(status_code=403, detail="Solo puedes programar sesiones de tus formaciones")
    try:
        starts_at, ends_at = validate_session_schedule(request.starts_at, request.ends_at, request.timezone)
    except ValueError as error:
        raise HTTPException(status_code=422, detail=str(error)) from error
    participants = await validate_session_participants(
        request.participant_user_ids, content, current_user,
    )
    session = SynchronousSession(
        content_id=request.content_id,
        title=request.title.strip(),
        description=request.description,
        starts_at=starts_at,
        ends_at=ends_at,
        timezone=request.timezone,
        participant_user_ids=participants,
        join_window_minutes=request.join_window_minutes,
        created_by=current_user.id,
    )
    payload = mongo_serialize(session)
    payload["status"] = session.status.value
    if request.meeting_url:
        try:
            payload.update(meeting_url_storage_fields(request.meeting_url))
        except RuntimeError as error:
            raise HTTPException(status_code=503, detail=str(error)) from error
    await db.synchronous_sessions.insert_one(payload)
    await enqueue_event("SessionScheduled", "SynchronousSession", session.id, {
        "session_id": session.id, "starts_at": starts_at,
        "timezone": session.timezone, "user_ids": participants,
    }, f"session:{session.id}:scheduled:{starts_at.isoformat()}")
    await log_activity(current_user, "Creación de sesión", {
        "id": session.id, "name": session.title,
    }, {"content_id": session.content_id, "participant_count": len(participants)})
    return await build_session_view(payload, current_user)


@api_router.get("/sessions/eligible-users", response_model=List[User])
async def session_eligible_users(content_id: str, current_user: User = Depends(get_current_user)):
    content = await db.training_contents.find_one({"id": content_id}, {"_id": 0})
    if not content:
        raise HTTPException(status_code=404, detail="Formacion no encontrada")
    if current_user.user_type not in {UserType.ADMIN, UserType.JUNTA_DIRECTIVA, UserType.FORMADOR}:
        raise HTTPException(status_code=403, detail="Sin permisos para convocar participantes")
    if current_user.user_type == UserType.FORMADOR and content.get("created_by") != current_user.id:
        raise HTTPException(status_code=403, detail="Solo puedes gestionar participantes de tus formaciones")
    if current_user.user_type in {UserType.ADMIN, UserType.JUNTA_DIRECTIVA}:
        query: Dict[str, Any] = {"is_active": True, "user_type": {"$ne": UserType.ADMIN.value}}
    else:
        enrollment_ids = [item["user_id"] for item in await db.enrollments.find({
            "content_id": content_id, "status": {"$ne": EnrollmentStatus.CANCELLED.value},
        }, {"user_id": 1}).to_list(10000)]
        query = {"id": {"$in": enrollment_ids}, "is_active": True}
    users = await db.users.find(query, {"_id": 0, "password_hash": 0}).sort("name", 1).to_list(10000)
    return [User(**item) for item in users]


@api_router.get("/sessions/{session_id}", response_model=SynchronousSessionView)
async def get_synchronous_session(
    session_id: str, current_user: User = Depends(get_current_user),
):
    return await build_session_view(await require_session_access(current_user, session_id), current_user)


@api_router.patch("/sessions/{session_id}", response_model=SynchronousSessionView)
async def update_synchronous_session(
    session_id: str,
    request: SessionUpdateRequest,
    current_user: User = Depends(get_current_user),
):
    session = await db.synchronous_sessions.find_one({"id": session_id})
    if not session:
        raise HTTPException(status_code=404, detail="Sesión no encontrada")
    if not await can_manage_session(current_user, session):
        raise HTTPException(status_code=403, detail="No puedes editar esta sesión")

    changes = request.model_dump(exclude_unset=True, mode="python")
    meeting_url = changes.pop("meeting_url", None)
    timezone_name = changes.get("timezone", session.get("timezone", DEFAULT_TIMEZONE))
    starts_at = changes.get("starts_at", session["starts_at"])
    ends_at = changes.get("ends_at", session["ends_at"])
    try:
        normalized_start, normalized_end = validate_session_schedule(starts_at, ends_at, timezone_name)
    except ValueError as error:
        raise HTTPException(status_code=422, detail=str(error)) from error
    if "starts_at" in changes or "ends_at" in changes:
        changes["starts_at"] = normalized_start
        changes["ends_at"] = normalized_end

    if request.status is not None:
        current_status = SynchronousSessionStatus(
            session.get("status", SynchronousSessionStatus.SCHEDULED.value)
        )
        allowed_transitions = {
            SynchronousSessionStatus.DRAFT: {
                SynchronousSessionStatus.SCHEDULED, SynchronousSessionStatus.CANCELLED,
            },
            SynchronousSessionStatus.SCHEDULED: {
                SynchronousSessionStatus.LIVE, SynchronousSessionStatus.ENDED,
                SynchronousSessionStatus.CANCELLED,
            },
            SynchronousSessionStatus.LIVE: {
                SynchronousSessionStatus.ENDED, SynchronousSessionStatus.CANCELLED,
            },
            SynchronousSessionStatus.ENDED: set(),
            SynchronousSessionStatus.CANCELLED: set(),
        }
        if request.status != current_status and request.status not in allowed_transitions[current_status]:
            raise HTTPException(status_code=409, detail="Transición de estado no permitida")
        changes["status"] = request.status.value

    if "title" in changes and changes["title"]:
        changes["title"] = changes["title"].strip()
    changes["updated_at"] = datetime.now(timezone.utc)
    update: Dict[str, Any] = {"$set": changes}
    if "meeting_url" in request.model_fields_set:
        if meeting_url:
            try:
                update["$set"].update(meeting_url_storage_fields(meeting_url))
            except RuntimeError as error:
                raise HTTPException(status_code=503, detail=str(error)) from error
            update["$unset"] = {"meeting_url": ""}
        else:
            update["$unset"] = {
                "meeting_url": "", "meeting_url_ciphertext": "", "meeting_url_hash": "",
            }
    updated = await db.synchronous_sessions.find_one_and_update(
        {"_id": session["_id"]}, update, return_document=ReturnDocument.AFTER,
    )
    changed_schedule = any(field in request.model_fields_set for field in {"starts_at", "ends_at", "timezone"})
    event_type = "SessionCancelled" if changes.get("status") == SynchronousSessionStatus.CANCELLED.value else "SessionRescheduled" if changed_schedule else "SessionUpdated"
    await enqueue_event(event_type, "SynchronousSession", session_id, {
        "session_id": session_id, "starts_at": updated["starts_at"],
        "timezone": updated.get("timezone", DEFAULT_TIMEZONE),
        "user_ids": updated.get("participant_user_ids", []),
    }, f"session:{session_id}:{event_type}:{updated['updated_at'].isoformat()}")
    await log_activity(current_user, "Actualización de sesión", {
        "id": session_id, "name": session.get("title", "Sesión")
    }, {"changed_fields": sorted(request.model_fields_set)})
    return await build_session_view(updated, current_user)


@api_router.put("/sessions/{session_id}/participants", response_model=SynchronousSessionView)
async def update_session_participants(
    session_id: str,
    request: SessionParticipantsRequest,
    current_user: User = Depends(get_current_user),
):
    session = await db.synchronous_sessions.find_one({"id": session_id})
    if not session:
        raise HTTPException(status_code=404, detail="Sesión no encontrada")
    if not await can_manage_session(current_user, session):
        raise HTTPException(status_code=403, detail="No puedes editar esta sesión")
    content = await db.training_contents.find_one({"id": session["content_id"]}, {"_id": 0})
    if not content:
        raise HTTPException(status_code=409, detail="La formación asociada ya no existe")
    participants = await validate_session_participants(request.user_ids, content, current_user)
    updated = await db.synchronous_sessions.find_one_and_update(
        {"_id": session["_id"]},
        {"$set": {"participant_user_ids": participants, "updated_at": datetime.now(timezone.utc)}},
        return_document=ReturnDocument.AFTER,
    )
    await enqueue_event("SessionParticipantsUpdated", "SynchronousSession", session_id, {
        "session_id": session_id, "starts_at": updated["starts_at"],
        "timezone": updated.get("timezone", DEFAULT_TIMEZONE), "user_ids": participants,
    }, f"session:{session_id}:participants:{updated['updated_at'].isoformat()}")
    await log_activity(current_user, "Actualización de participantes", {
        "id": session_id, "name": session.get("title", "Sesión")
    }, {"participant_count": len(participants)})
    return await build_session_view(updated, current_user)


@api_router.get("/sessions/{session_id}/participants", response_model=List[User])
async def get_session_participants(session_id: str, current_user: User = Depends(get_current_user)):
    session = await db.synchronous_sessions.find_one({"id": session_id})
    if not session:
        raise HTTPException(status_code=404, detail="Sesion no encontrada")
    if not await can_manage_session(current_user, session):
        raise HTTPException(status_code=403, detail="No puedes consultar los participantes")
    user_ids = session.get("participant_user_ids", [])
    if not user_ids:
        return []
    users = await db.users.find({"id": {"$in": user_ids}}, {"_id": 0, "password_hash": 0}).sort("name", 1).to_list(10000)
    return [User(**item) for item in users]


@api_router.delete("/sessions/{session_id}", status_code=204)
async def cancel_synchronous_session(
    session_id: str, current_user: User = Depends(get_current_user),
):
    session = await db.synchronous_sessions.find_one({"id": session_id})
    if not session:
        raise HTTPException(status_code=404, detail="Sesión no encontrada")
    if not await can_manage_session(current_user, session):
        raise HTTPException(status_code=403, detail="No puedes cancelar esta sesión")
    if session.get("status") == SynchronousSessionStatus.ENDED.value:
        raise HTTPException(status_code=409, detail="Una sesión finalizada no puede cancelarse")
    await db.synchronous_sessions.update_one({"_id": session["_id"]}, {"$set": {
        "status": SynchronousSessionStatus.CANCELLED.value,
        "updated_at": datetime.now(timezone.utc),
    }})
    await enqueue_event("SessionCancelled", "SynchronousSession", session_id, {
        "session_id": session_id, "starts_at": session["starts_at"],
        "timezone": session.get("timezone", DEFAULT_TIMEZONE),
        "user_ids": session.get("participant_user_ids", []),
    }, f"session:{session_id}:cancelled")
    await log_activity(current_user, "Cancelación de sesión", {
        "id": session_id, "name": session.get("title", "Sesión")
    })


@api_router.post("/sessions/{session_id}/join")
async def join_synchronous_session(
    session_id: str, current_user: User = Depends(get_current_user),
):
    session = await require_session_access(current_user, session_id)
    allowed, reason = session_join_state(session)
    if not allowed:
        raise HTTPException(status_code=403, detail=reason)
    meeting_url = read_meeting_url(session)
    if not meeting_url:
        raise HTTPException(status_code=503, detail="El enlace de acceso no está disponible")
    await log_activity(current_user, "Acceso a sesión", {
        "id": session_id, "name": session.get("title", "Sesión")
    }, {"content_id": session["content_id"]})
    return RedirectResponse(
        meeting_url,
        status_code=303,
        headers={"Cache-Control": "no-store", "Referrer-Policy": "no-referrer"},
    )


@api_router.get("/notifications")
async def list_notifications(current_user: User = Depends(get_current_user)):
    items = await db.notifications.find(
        {
            "user_id": current_user.id,
            "in_app_visible": {"$ne": False},
            "scheduled_at": {"$lte": datetime.now(timezone.utc)},
            "status": {"$ne": "cancelled"},
        }, {"_id": 0}
    ).sort("created_at", -1).limit(100).to_list(100)
    return items


@api_router.post("/notifications/{notification_id}/read")
async def read_notification(notification_id: str, current_user: User = Depends(get_current_user)):
    result = await db.notifications.update_one(
        {"id": notification_id, "user_id": current_user.id},
        {"$set": {"read_at": datetime.now(timezone.utc)}},
    )
    if not result.matched_count:
        raise HTTPException(status_code=404, detail="Notificacion no encontrada")
    return {"success": True}


@api_router.post("/notifications/read-all")
async def read_all_notifications(current_user: User = Depends(get_current_user)):
    result = await db.notifications.update_many(
        {"user_id": current_user.id, "read_at": None}, {"$set": {"read_at": datetime.now(timezone.utc)}},
    )
    return {"updated": result.modified_count}


@api_router.get("/notification-preferences")
async def get_notification_preferences(current_user: User = Depends(get_current_user)):
    stored = await db.notification_preferences.find_one({"user_id": current_user.id}, {"_id": 0})
    return {
        "user_id": current_user.id,
        **NotificationPreferencesRequest().model_dump(),
        **(stored or {}),
    }


@api_router.put("/notification-preferences")
async def update_notification_preferences(request: NotificationPreferencesRequest, current_user: User = Depends(get_current_user)):
    payload = request.model_dump(mode="python")
    payload["updated_at"] = datetime.now(timezone.utc)
    await db.notification_preferences.update_one({"user_id": current_user.id}, {"$set": payload, "$setOnInsert": {"user_id": current_user.id}}, upsert=True)
    return {"user_id": current_user.id, **request.model_dump()}


@api_router.get("/integrations/telegram")
async def telegram_status(current_user: User = Depends(get_current_user)):
    binding = await db.telegram_bindings.find_one({"user_id": current_user.id, "active": True}, {"_id": 0, "chat_id": 0, "telegram_user_id": 0})
    return {"configured": bool(os.getenv("TELEGRAM_BOT_TOKEN") and os.getenv("TELEGRAM_BOT_USERNAME")), "linked": bool(binding), "linked_at": binding.get("linked_at") if binding else None}


@api_router.post("/integrations/telegram/link")
async def create_telegram_link(current_user: User = Depends(get_current_user)):
    username = os.getenv("TELEGRAM_BOT_USERNAME")
    if not username or not os.getenv("TELEGRAM_BOT_TOKEN"):
        raise HTTPException(status_code=503, detail="Telegram no esta configurado")
    raw_token = secrets.token_urlsafe(32)
    now = datetime.now(timezone.utc)
    await db.telegram_link_tokens.update_many({"user_id": current_user.id, "used_at": None}, {"$set": {"used_at": now}})
    await db.telegram_link_tokens.insert_one({
        "id": str(uuid.uuid4()), "user_id": current_user.id, "token_hash": token_digest(raw_token),
        "created_at": now, "expires_at": now + timedelta(minutes=10), "used_at": None,
    })
    return {"url": telegram_link(username, raw_token), "expires_in_minutes": 10}


@api_router.delete("/integrations/telegram", status_code=204)
async def unlink_telegram(current_user: User = Depends(get_current_user)):
    await db.telegram_bindings.update_many({"user_id": current_user.id}, {"$set": {"active": False, "unlinked_at": datetime.now(timezone.utc)}})


@api_router.post("/webhooks/telegram")
async def telegram_webhook(request: Request):
    configured_secret = os.getenv("TELEGRAM_WEBHOOK_SECRET")
    supplied_secret = request.headers.get("X-Telegram-Bot-Api-Secret-Token", "")
    if not configured_secret or not hmac.compare_digest(configured_secret, supplied_secret):
        raise HTTPException(status_code=401, detail="Firma de webhook no valida")
    payload = await request.json()
    message = payload.get("message") or {}
    chat = message.get("chat") or {}
    sender = message.get("from") or {}
    if chat.get("type") != "private" or not isinstance(chat.get("id"), int) or not isinstance(sender.get("id"), int):
        return {"ok": True}
    text = str(message.get("text", "")).strip()
    if text == "/stop":
        await db.telegram_bindings.update_many({"telegram_user_id": sender["id"]}, {"$set": {"active": False, "unlinked_at": datetime.now(timezone.utc)}})
        return {"ok": True}
    if not text.startswith("/start "):
        return {"ok": True}
    raw_token = text.split(maxsplit=1)[1]
    now = datetime.now(timezone.utc)
    link_token = await db.telegram_link_tokens.find_one_and_update(
        {"token_hash": token_digest(raw_token), "used_at": None, "expires_at": {"$gt": now}},
        {"$set": {"used_at": now}}, return_document=ReturnDocument.AFTER,
    )
    if not link_token:
        return {"ok": True}
    await db.telegram_bindings.update_many(
        {"$or": [{"user_id": link_token["user_id"]}, {"telegram_user_id": sender["id"]}]},
        {"$set": {"active": False, "unlinked_at": now}},
    )
    await db.telegram_bindings.insert_one({
        "id": str(uuid.uuid4()), "user_id": link_token["user_id"],
        "telegram_user_id": sender["id"], "chat_id": chat["id"], "active": True, "linked_at": now,
    })
    return {"ok": True}


async def session_query_for_user(user: User) -> Dict[str, Any]:
    if user.user_type in {UserType.ADMIN, UserType.JUNTA_DIRECTIVA}:
        return {}
    await ensure_legacy_enrollments_for_user(user)
    content_ids = [item["content_id"] for item in await db.enrollments.find({
        "user_id": user.id, "status": {"$ne": EnrollmentStatus.CANCELLED.value},
    }, {"content_id": 1}).to_list(5000)]
    return {"$or": [
        {"participant_user_ids": user.id},
        {"$and": [{"$or": [{"participant_user_ids": {"$size": 0}}, {"participant_user_ids": {"$exists": False}}]}, {"content_id": {"$in": content_ids}}]},
    ]}


@api_router.post("/integrations/calendar/feed")
async def create_calendar_feed(current_user: User = Depends(get_current_user)):
    raw_token = secrets.token_urlsafe(40)
    now = datetime.now(timezone.utc)
    await db.calendar_feed_tokens.update_many({"user_id": current_user.id, "revoked_at": None}, {"$set": {"revoked_at": now}})
    await db.calendar_feed_tokens.insert_one({"id": str(uuid.uuid4()), "user_id": current_user.id, "token_hash": token_digest(raw_token), "created_at": now, "revoked_at": None})
    return {"url": f"{PUBLIC_API_URL.rstrip('/')}/api/calendar/feed/{raw_token}.ics"}


@api_router.delete("/integrations/calendar/feed", status_code=204)
async def revoke_calendar_feed(current_user: User = Depends(get_current_user)):
    await db.calendar_feed_tokens.update_many({"user_id": current_user.id, "revoked_at": None}, {"$set": {"revoked_at": datetime.now(timezone.utc)}})


@api_router.get("/calendar/feed/{token}.ics")
async def calendar_feed(token: str):
    stored = await db.calendar_feed_tokens.find_one({"token_hash": token_digest(token), "revoked_at": None})
    if not stored:
        raise HTTPException(status_code=404, detail="Calendario no encontrado")
    user_data = await db.users.find_one({"id": stored["user_id"], "is_active": True}, {"_id": 0, "password_hash": 0})
    if not user_data:
        raise HTTPException(status_code=404, detail="Calendario no encontrado")
    sessions = await db.synchronous_sessions.find(await session_query_for_user(User(**user_data)), {"_id": 0}).sort("starts_at", 1).to_list(5000)
    return PlainTextResponse(render_ics(sessions, PUBLIC_APP_URL), media_type="text/calendar; charset=utf-8", headers={"Cache-Control": "private, max-age=300", "Content-Disposition": "inline; filename=formacion-ritsi.ics"})


async def session_attendee_ids(session: Dict[str, Any]) -> List[str]:
    if session.get("participant_user_ids"):
        return session["participant_user_ids"]
    items = await db.enrollments.find({"content_id": session["content_id"], "status": {"$ne": EnrollmentStatus.CANCELLED.value}}, {"user_id": 1}).to_list(10000)
    return [item["user_id"] for item in items]


@api_router.get("/sessions/{session_id}/attendance")
async def get_session_attendance(session_id: str, current_user: User = Depends(get_current_user)):
    session = await db.synchronous_sessions.find_one({"id": session_id})
    if not session:
        raise HTTPException(status_code=404, detail="Sesion no encontrada")
    if not await can_manage_session(current_user, session):
        raise HTTPException(status_code=403, detail="No puedes consultar esta asistencia")
    attendee_ids = await session_attendee_ids(session)
    users = await db.users.find({"id": {"$in": attendee_ids}}, {"_id": 0, "password_hash": 0}).to_list(10000)
    records = await db.session_attendance.find({"session_id": session_id}, {"_id": 0}).to_list(10000)
    by_user = {item["user_id"]: item for item in records}
    return [{"user": User(**user), "attendance": by_user.get(user["id"])} for user in users]


@api_router.put("/sessions/{session_id}/attendance")
async def save_session_attendance(session_id: str, request: AttendanceRequest, current_user: User = Depends(get_current_user)):
    session = await db.synchronous_sessions.find_one({"id": session_id})
    if not session:
        raise HTTPException(status_code=404, detail="Sesion no encontrada")
    if not await can_manage_session(current_user, session):
        raise HTTPException(status_code=403, detail="No puedes registrar esta asistencia")
    allowed_ids = set(await session_attendee_ids(session))
    if any(entry.user_id not in allowed_ids for entry in request.entries):
        raise HTTPException(status_code=422, detail="La asistencia contiene personas no convocadas")
    now = datetime.now(timezone.utc)
    for entry in request.entries:
        await db.session_attendance.update_one(
            {"session_id": session_id, "user_id": entry.user_id},
            {"$set": {**entry.model_dump(), "session_id": session_id, "source": "manual", "updated_by": current_user.id, "updated_at": now}, "$setOnInsert": {"id": str(uuid.uuid4()), "created_at": now}}, upsert=True,
        )
    return {"updated": len(request.entries)}


@api_router.get("/reports/overview")
async def reports_overview(current_user: User = Depends(get_current_user)):
    if current_user.user_type not in {UserType.ADMIN, UserType.JUNTA_DIRECTIVA, UserType.UNIVERSIDAD, UserType.FORMADOR}:
        raise HTTPException(status_code=403, detail="Sin permisos para consultar informes")
    enrollment_query: Dict[str, Any] = {}
    session_query: Dict[str, Any] = {}
    if current_user.user_type == UserType.UNIVERSIDAD:
        ids = [item["id"] for item in await db.users.find({"university_id": current_user.university_id}, {"id": 1}).to_list(10000)]
        enrollment_query = {"user_id": {"$in": ids}}
    elif current_user.user_type == UserType.FORMADOR:
        content_ids = [item["id"] for item in await db.training_contents.find({"created_by": current_user.id}, {"id": 1}).to_list(10000)]
        enrollment_query = {"content_id": {"$in": content_ids}}
        session_query = {"content_id": {"$in": content_ids}}
    statuses = {}
    for value in EnrollmentStatus:
        statuses[value.value] = await db.enrollments.count_documents({**enrollment_query, "status": value.value})
    session_ids = [item["id"] for item in await db.synchronous_sessions.find(session_query, {"id": 1}).to_list(10000)]
    attendance = await db.session_attendance.aggregate([
        {"$match": {"session_id": {"$in": session_ids}}}, {"$group": {"_id": "$status", "count": {"$sum": 1}}},
    ]).to_list(20)
    return {"enrollments": statuses, "sessions": len(session_ids), "attendance": {item["_id"]: item["count"] for item in attendance}}


@api_router.get("/operations/integrations")
async def integration_operations(current_user: User = Depends(get_current_user)):
    if current_user.user_type != UserType.ADMIN:
        raise HTTPException(status_code=403, detail="Solo administracion puede consultar operaciones")
    failed_events = await db.outbox_events.find({"status": "failed"}, {"_id": 0, "id": 1, "event_type": 1, "aggregate_id": 1, "last_error": 1, "attempts": 1}).sort("occurred_at", -1).limit(20).to_list(20)
    failed_calendar = await db.calendar_event_links.find({"status": "failed"}, {"_id": 0, "session_id": 1, "last_error": 1, "updated_at": 1}).sort("updated_at", -1).limit(20).to_list(20)
    return {
        "outbox": {
            status: await db.outbox_events.count_documents({"status": status})
            for status in ("pending", "processing", "processed", "failed")
        },
        "telegram": {
            "configured": bool(os.getenv("TELEGRAM_BOT_TOKEN") and os.getenv("TELEGRAM_WEBHOOK_SECRET")),
            "active_bindings": await db.telegram_bindings.count_documents({"active": True}),
        },
        "google_calendar": {
            "configured": all(os.getenv(name) for name in ("GOOGLE_CALENDAR_ID", "GOOGLE_CLIENT_ID", "GOOGLE_CLIENT_SECRET", "GOOGLE_REFRESH_TOKEN")),
            "failed": await db.calendar_event_links.count_documents({"status": "failed"}),
        },
        "failed_events": failed_events,
        "failed_calendar_events": failed_calendar,
    }


@api_router.post("/operations/outbox/{event_id}/retry")
async def retry_outbox_event(event_id: str, current_user: User = Depends(get_current_user)):
    if current_user.user_type != UserType.ADMIN:
        raise HTTPException(status_code=403, detail="Solo administracion puede reintentar eventos")
    result = await db.outbox_events.update_one(
        {"id": event_id, "status": "failed"},
        {"$set": {"status": "pending", "attempts": 0, "available_at": datetime.now(timezone.utc), "last_error": None}},
    )
    if not result.matched_count:
        raise HTTPException(status_code=404, detail="Evento fallido no encontrado")
    return {"success": True}


@api_router.post("/operations/calendar/{session_id}/resync")
async def resync_calendar_event(session_id: str, current_user: User = Depends(get_current_user)):
    if current_user.user_type != UserType.ADMIN:
        raise HTTPException(status_code=403, detail="Solo administracion puede reconciliar calendarios")
    session = await db.synchronous_sessions.find_one({"id": session_id}, {"_id": 0})
    if not session:
        raise HTTPException(status_code=404, detail="Sesion no encontrada")
    await enqueue_event("SessionUpdated", "SynchronousSession", session_id, {
        "session_id": session_id, "starts_at": session["starts_at"],
        "user_ids": session.get("participant_user_ids", []),
    }, f"session:{session_id}:manual-resync:{uuid.uuid4()}")
    return {"queued": True}


@api_router.get("/activity-log", response_model=List[ActivityLog])
async def activity_log(current_user: User = Depends(get_current_user)):
    if current_user.user_type != UserType.ADMIN:
        raise HTTPException(status_code=403, detail="Sin permisos")
    return [ActivityLog(**item) for item in await db.activity_logs.find({}, {"_id": 0}).sort("timestamp", -1).to_list(2000)]


async def migrate_legacy_security_and_session_records() -> None:
    legacy_auth_sessions = await db.user_sessions.find({}).to_list(10000)
    for item in legacy_auth_sessions:
        changes: Dict[str, Any] = {}
        removals: Dict[str, str] = {}
        if item.get("session_token"):
            if not item.get("session_token_hash"):
                changes["session_token_hash"] = token_digest(item["session_token"])
            removals["session_token"] = ""
        for field in ("created_at", "expires_at"):
            if isinstance(item.get(field), str):
                try:
                    changes[field] = parse_stored_utc_datetime(item[field])
                except ValueError:
                    logger.warning("Sesion de autenticacion con %s invalido: %s", field, item.get("id"))
        update: Dict[str, Any] = {}
        if changes:
            update["$set"] = changes
        if removals:
            update["$unset"] = removals
        if update:
            await db.user_sessions.update_one({"_id": item["_id"]}, update)

    legacy_sync_sessions = await db.synchronous_sessions.find({}).to_list(10000)
    for item in legacy_sync_sessions:
        changes = {}
        removals = {}
        for field in ("starts_at", "ends_at", "created_at", "updated_at"):
            if isinstance(item.get(field), str):
                try:
                    changes[field] = parse_stored_utc_datetime(item[field])
                except ValueError:
                    logger.warning("Sesion sincrona con %s invalido: %s", field, item.get("id"))
        if item.get("meeting_url") and MEETING_URL_ENCRYPTION_KEY:
            changes.update(meeting_url_storage_fields(item["meeting_url"]))
            removals["meeting_url"] = ""
        update = {}
        if changes:
            update["$set"] = changes
        if removals:
            update["$unset"] = removals
        if update:
            await db.synchronous_sessions.update_one({"_id": item["_id"]}, update)
    await db.activity_logs.update_many(
        {"details.password_hash": {"$exists": True}},
        {"$unset": {"details.password_hash": ""}},
    )


async def encrypt_pending_legacy_meeting_urls() -> None:
    if not MEETING_URL_ENCRYPTION_KEY:
        return
    legacy_sessions = await db.synchronous_sessions.find({
        "meeting_url": {"$type": "string"},
    }).to_list(10000)
    for item in legacy_sessions:
        meeting_url = item.get("meeting_url")
        if not meeting_url:
            continue
        await db.synchronous_sessions.update_one(
            {"_id": item["_id"]},
            {
                "$set": meeting_url_storage_fields(meeting_url),
                "$unset": {"meeting_url": ""},
            },
        )


async def migrate_legacy_assignments_and_progress() -> None:
    assignments = await db.content_assignments.find({}).to_list(10000)
    for assignment in assignments:
        for user_id in set(assignment.get("assigned_to_user_ids", [])):
            await upsert_enrollment(
                user_id,
                assignment["content_id"],
                assignment["assigned_by"],
                EnrollmentSource.ASSIGNMENT,
                assignment["id"],
            )
    progress_items = await db.user_progress.find({}).to_list(10000)
    content_cache: Dict[str, Optional[Dict[str, Any]]] = {}
    for progress in progress_items:
        content_id = progress.get("content_id")
        if not content_id:
            continue
        if content_id not in content_cache:
            content_cache[content_id] = await db.training_contents.find_one({"id": content_id}, {"_id": 0})
        content = content_cache[content_id]
        if content:
            await db.user_progress.update_one(
                {"_id": progress["_id"]},
                {"$set": {"completed": calculate_progress_completion(content, progress)}},
            )


async def run_phase_one_migrations() -> None:
    migration_id = "2026-08-phase-1-security-enrollments-sessions"
    if await db.schema_migrations.find_one({"id": migration_id}, {"id": 1}):
        return
    await migrate_legacy_security_and_session_records()
    await migrate_legacy_assignments_and_progress()
    await db.schema_migrations.update_one(
        {"id": migration_id},
        {"$setOnInsert": {
            "id": migration_id,
            "applied_at": datetime.now(timezone.utc),
        }},
        upsert=True,
    )


async def run_phase_two_migrations() -> None:
    migration_id = "2026-08-phase-2-notifications-certificates"
    if await db.schema_migrations.find_one({"id": migration_id}, {"id": 1}):
        return
    completed_progress = await db.user_progress.find({"completed": True}, {"user_id": 1, "content_id": 1, "last_updated": 1}).to_list(100000)
    for progress in completed_progress:
        user = await db.users.find_one({"id": progress.get("user_id")}, {"name": 1})
        content = await db.training_contents.find_one({"id": progress.get("content_id")}, {"title": 1})
        if not user or not content:
            continue
        await db.certificates.update_one(
            {"user_id": progress["user_id"], "content_id": progress["content_id"]},
            {"$setOnInsert": {
                "id": str(uuid.uuid4()), "verification_code": secrets.token_urlsafe(16),
                "user_id": progress["user_id"], "user_name": user["name"],
                "content_id": progress["content_id"], "content_title": content["title"],
                "issued_at": progress.get("last_updated") or datetime.now(timezone.utc), "status": "valid",
            }}, upsert=True,
        )
    await db.schema_migrations.update_one(
        {"id": migration_id}, {"$setOnInsert": {"id": migration_id, "applied_at": datetime.now(timezone.utc)}}, upsert=True,
    )


async def initialize_database():
    allowed_roles = [item.value for item in UserType]
    await db.users.update_many({"user_type": {"$nin": allowed_roles}}, {"$set": {"user_type": UserType.FORMADOR.value, "board_position": None}})
    await db.users.update_many({"vocalia_ids": {"$exists": False}}, {"$set": {"vocalia_ids": []}})
    legacy_contents = await db.training_contents.find({"files.google_drive_url": {"$exists": True}}).to_list(10000)
    for content in legacy_contents:
        files = []
        for item in content.get("files", []):
            resource = dict(item)
            resource["url"] = resource.pop("google_drive_url", resource.get("url", ""))
            resource["file_type"] = resource.get("file_type", FileType.LINK.value)
            files.append(resource)
        await db.training_contents.update_one({"_id": content["_id"]}, {"$set": {"files": files}})
    await run_phase_one_migrations()
    await run_phase_two_migrations()
    await encrypt_pending_legacy_meeting_urls()
    await db.schema_migrations.create_index("id", unique=True)
    await db.users.create_index("email", unique=True)
    await db.user_sessions.create_index("session_token_hash", unique=True, sparse=True)
    await db.user_sessions.create_index("expires_at", expireAfterSeconds=0)
    await db.training_contents.create_index("id", unique=True)
    await db.training_contents.create_index("external_id", unique=True, sparse=True)
    await db.content_assignments.create_index("id", unique=True)
    await db.content_assignments.create_index([("content_id", 1), ("assigned_by", 1)])
    await db.content_assignments.create_index("assigned_to_user_ids")
    await db.enrollments.create_index([("content_id", 1), ("user_id", 1)], unique=True)
    await db.enrollments.create_index([("user_id", 1), ("status", 1), ("created_at", -1)])
    await db.user_progress.create_index([("user_id", 1), ("content_id", 1)], unique=True)
    await db.synchronous_sessions.create_index("id", unique=True)
    await db.synchronous_sessions.create_index([("content_id", 1), ("starts_at", 1)])
    await db.synchronous_sessions.create_index([("participant_user_ids", 1), ("starts_at", 1)])
    await db.synchronous_sessions.create_index("meeting_url_hash", sparse=True)
    await db.outbox_events.create_index("dedupe_key", unique=True)
    await db.outbox_events.create_index([("status", 1), ("available_at", 1), ("occurred_at", 1)])
    await db.notifications.create_index("dedupe_key", unique=True)
    await db.notifications.create_index([("user_id", 1), ("created_at", -1)])
    await db.notifications.create_index([("status", 1), ("scheduled_at", 1)])
    await db.notification_deliveries.create_index("dedupe_key", unique=True)
    await db.notification_deliveries.create_index([("status", 1), ("next_attempt_at", 1)])
    await db.notification_deliveries.create_index([("notification_id", 1), ("channel", 1)], unique=True)
    await db.notification_preferences.create_index("user_id", unique=True)
    await db.telegram_link_tokens.create_index("token_hash", unique=True)
    await db.telegram_link_tokens.create_index("expires_at", expireAfterSeconds=0)
    await db.telegram_bindings.create_index("user_id", unique=True, partialFilterExpression={"active": True})
    await db.telegram_bindings.create_index("telegram_user_id", unique=True, partialFilterExpression={"active": True})
    await db.calendar_feed_tokens.create_index("token_hash", unique=True)
    await db.calendar_event_links.create_index([("session_id", 1), ("calendar_id", 1)], unique=True)
    await db.password_reset_tokens.create_index("token_hash", unique=True)
    await db.password_reset_tokens.create_index("expires_at", expireAfterSeconds=0)
    await db.auth_rate_limits.create_index("key", unique=True)
    await db.auth_rate_limits.create_index("expires_at", expireAfterSeconds=0)
    await db.session_attendance.create_index([("session_id", 1), ("user_id", 1)], unique=True)
    await db.assessment_attempts.create_index([("user_id", 1), ("content_id", 1), ("quiz_id", 1), ("attempt_number", 1)], unique=True)
    await db.learning_paths.create_index("id", unique=True)
    await db.learning_path_assignments.create_index("id", unique=True)
    await db.learning_path_assignments.create_index("user_ids")
    await db.certificates.create_index([("user_id", 1), ("content_id", 1)], unique=True)
    await db.certificates.create_index("verification_code", unique=True)
    await db.privacy_requests.create_index([("user_id", 1), ("status", 1)])
    await db.activity_logs.create_index("timestamp")
    await db.vocalias.create_index("name", unique=True)
    await db.universities.create_index("source_key", unique=True, sparse=True)
    if os.getenv("AUTO_IMPORT_UNIVERSITIES", "true").lower() == "true" and await db.universities.count_documents({}) == 0:
        try:
            await import_ritsi_universities()
        except Exception as error:
            logger.warning("No se pudo realizar la importación inicial de universidades: %s", error)
    if os.getenv("AUTO_IMPORT_FORMATIONS", "true").lower() == "true" and await db.training_contents.count_documents({}) == 0:
        try:
            await import_formations(SHEET_ID, "sheet-import", True)
        except Exception as error:
            logger.warning("No se pudo realizar la importación inicial del catálogo: %s", error)


@asynccontextmanager
async def lifespan(application: FastAPI):
    await initialize_database()
    yield
    client.close()


app = FastAPI(title="Plataforma Formativa RITSI", version="2.0.0", lifespan=lifespan)
app.include_router(api_router)


@app.middleware("http")
async def protect_cookie_authenticated_mutations(request: Request, call_next):
    """Reject browser cross-site mutations before they reach authenticated routes."""

    if request.method in UNSAFE_HTTP_METHODS and request.cookies.get("session_token"):
        origin = request.headers.get("origin")
        fetch_site = request.headers.get("sec-fetch-site")
        if not browser_mutation_origin_is_trusted(origin, fetch_site):
            return JSONResponse(
                {"detail": "Origen no autorizado para esta operacion"},
                status_code=403,
            )
    return await call_next(request)


app.add_middleware(CORSMiddleware, allow_credentials=True,
    allow_origins=sorted(ALLOWED_ORIGINS),
    allow_methods=["*"], allow_headers=["*"])
